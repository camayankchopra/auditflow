"""Audit Sampling Tool - builds a sampling workpaper (Excel) from a population listing.

Methods: MUS (monetary unit sampling), Random (seeded), Systematic, Haphazard (auditor selects).
Sample size is driven by assessed risk, test-of-controls result, analytical procedures and
expected misstatement, using the editable tables on the 'Parameters' tab.

IMPORTANT: the default tables are ILLUSTRATIVE (Poisson-based, AICPA-style reliability factors).
Each firm has its own methodology - replace the Parameters tab values with your firm's tables.
"""
import bisect
import csv
import datetime as dt
import math
import random
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

# ----------------------------------------------------------------------------- parameters
IR_OPTS = ["Low", "Moderate", "High"]
TOC_OPTS = ["Not tested / not relied on", "Tested - effective", "Tested - ineffective"]
SAP_OPTS = ["None", "Limited", "Effective"]
RMM_ORDER = ["High", "Moderate", "Low"]
METHODS = ["MUS", "Random", "Systematic", "Haphazard"]
METHOD_NOTES = {
    "MUS": "Monetary unit sampling - every rupee/dollar has an equal chance; larger items more likely selected",
    "Random": "Simple random selection of items (seeded, reproducible)",
    "Systematic": "Every nth item from a random start (seeded, reproducible)",
    "Haphazard": "Auditor selects items without conscious bias - enter the Item IDs on the Sample tab",
}
ASSERTIONS = ["Existence", "Completeness", "Valuation / accuracy", "Rights & obligations",
              "Occurrence", "Cut-off", "Presentation & disclosure"]

# Risk of material misstatement: rows = inherent risk, columns = TOC_OPTS
RMM_MATRIX = {"Low": ["Low", "Low", "Moderate"],
              "Moderate": ["Moderate", "Low", "High"],
              "High": ["High", "Moderate", "High"]}
# Reliability (confidence) factor: rows = RMM, columns = SAP_OPTS  (Poisson, zero expected misstatement)
RF_TABLE = {"High": [3.00, 2.70, 2.31], "Moderate": [2.31, 2.00, 1.61], "Low": [1.61, 1.39, 1.21]}
EF_TABLE = {"High": 1.6, "Moderate": 1.5, "Low": 1.3}   # expansion factor for expected misstatement
SUPPORT_PCT = 0.5                                         # non-MUS: supports if total <= 50% of TM

MAX_ROWS = 50000


def derive(ir, toc, sig, sap):
    rmm = "High" if sig else RMM_MATRIX[ir][TOC_OPTS.index(toc)]
    sap_eff = "None" if sig else sap
    rf = RF_TABLE[rmm][SAP_OPTS.index(sap_eff)]
    return rmm, sap_eff, rf, EF_TABLE[rmm]


# ----------------------------------------------------------------------------- reading input
ID_SYN = {"itemid", "id", "ref", "reference", "invoiceno", "invoicenumber", "invoice", "docno", "documentno",
          "documentnumber", "docnumber", "voucherno", "number", "no", "entryno", "jeno", "transactionid", "item"}
DESC_SYN = {"description", "desc", "party", "customer", "vendor", "supplier", "name", "narration", "details",
            "particulars", "accountname", "partyname", "text"}
AMT_SYN = {"amount", "bookvalue", "balance", "value", "amountinr", "grossamount", "closingbalance", "outstanding",
           "carryingamount", "netbookvalue", "bookamount", "amt"}


def _norm(s):
    return re.sub(r"[^a-z0-9]", "", str(s or "").lower())


def _num(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "").replace("₹", "").replace("$", "")
    if s in ("", "-"):
        return 0.0
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    if s.lower().endswith("cr"):
        neg, s = True, s[:-2].strip()
    elif s.lower().endswith("dr"):
        s = s[:-2].strip()
    val = float(s)
    return -val if neg else val


def read_population(path):
    """Return list of dicts {id, desc, amount}. Raises ValueError with a readable message."""
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
            raw = [r for r in csv.reader(f)]
    else:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb["Population"] if "Population" in wb.sheetnames else wb[wb.sheetnames[0]]
        raw = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
    hdr_i = None
    for i, r in enumerate(raw[:25]):
        if any(_norm(c) in AMT_SYN for c in r):
            hdr_i = i
            break
    if hdr_i is None:
        raise ValueError("Could not find an amount column. Use a header such as 'Amount' or 'Book value'.")
    hdr = [_norm(c) for c in raw[hdr_i]]
    ci = next((j for j, h in enumerate(hdr) if h in ID_SYN), None)
    cd = next((j for j, h in enumerate(hdr) if h in DESC_SYN), None)
    ca = next(j for j, h in enumerate(hdr) if h in AMT_SYN)
    rows, bad, seen = [], [], {}
    for k, r in enumerate(raw[hdr_i + 1:], start=hdr_i + 2):
        r = list(r) + [None] * (len(hdr) - len(r))
        idv = r[ci] if ci is not None else None
        desc = r[cd] if cd is not None else None
        if all(c in (None, "") for c in r):
            continue
        txt = f"{idv or ''} {desc or ''}".lower()
        if re.match(r"\s*(grand\s+)?total", txt):
            continue
        try:
            amt = _num(r[ca])
        except ValueError:
            bad.append(k)
            continue
        idv = str(idv).strip() if idv not in (None, "") else f"Item {len(rows) + 1}"
        if isinstance(r[ci] if ci is not None else None, float) and idv.endswith(".0"):
            idv = idv[:-2]
        if idv in seen:                      # keep IDs unique so lookups are reliable
            seen[idv] += 1
            idv = f"{idv}-{seen[idv]}"
        else:
            seen[idv] = 1
        rows.append(dict(id=idv, desc=str(desc).strip() if desc not in (None, "") else "", amount=amt))
    if bad:
        raise ValueError(f"Non-numeric amount in file row(s): {', '.join(map(str, bad[:10]))}")
    if not rows:
        raise ValueError("No population items found.")
    if len(rows) > MAX_ROWS:
        raise ValueError(f"Population has {len(rows):,} items; the limit is {MAX_ROWS:,}.")
    return rows


# ----------------------------------------------------------------------------- selection
def select(rows, P):
    """Apply key-item rules, sample sizing and selection. Returns a result dict."""
    om = float(P["om"])
    pm = om * P["pm_pct"]
    tm = float(P["tm"]) if P.get("tm") not in (None, "") else pm
    em = float(P.get("em") or 0)
    key_thr = float(P["key_thr"]) if P.get("key_thr") not in (None, "") else tm
    rmm, sap_eff, rf, ef = derive(P["ir"], P["toc"], P["sig"], P["sap"])
    den = tm - em * ef
    if tm <= 0:
        raise ValueError("Tolerable misstatement must be greater than zero.")
    if den <= 0:
        raise ValueError("Expected misstatement is too high relative to tolerable misstatement "
                         "(tolerable misstatement - expected misstatement x expansion factor must be positive).")
    net = sum(r["amount"] for r in rows)
    sign = 1 if net >= 0 else -1
    key_ids = {x.strip().lower() for x in re.split(r"[,\n;]+", P.get("key_ids") or "") if x.strip()}
    cats = []
    for r in rows:
        a = r["amount"]
        if a == 0:
            cats.append("Zero value")
        elif r["id"].lower() in key_ids:
            cats.append("Key item - specified")
        elif abs(a) >= key_thr:
            cats.append("Key item - above threshold")
        elif P.get("key_neg") and a * sign < 0:
            cats.append("Key item - opposite sign")
        else:
            cats.append("Residual")
    method = P["method"]
    seed = P.get("seed")
    if seed in (None, ""):
        seed = random.SystemRandom().randint(100000, 999999)
    seed = int(seed)
    rng = random.Random(seed)

    def residual_idx():
        return [i for i, c in enumerate(cats) if c == "Residual"]

    def calc_n(R, cnt):
        if cnt == 0 or R <= 0:
            return 0
        return min(cnt, math.ceil(round(R * rf / den, 9)))

    # MUS: items >= sampling interval are selected with certainty ("top stratum") - iterate until stable
    if method == "MUS":
        while True:
            idx = residual_idx()
            R = sum(abs(rows[i]["amount"]) for i in idx)
            n = calc_n(R, len(idx))
            if n == 0:
                break
            interval = R / n
            top = [i for i in idx if abs(rows[i]["amount"]) >= interval]
            if not top or n >= len(idx):
                break
            for i in top:
                cats[i] = "Top stratum (MUS)"
    idx = residual_idx()
    R = sum(abs(rows[i]["amount"]) for i in idx)
    n = calc_n(R, len(idx))
    interval = (R / n) if (method == "MUS" and n) else None
    sample = []            # (row index, basis, hit point)
    start = None
    if n:
        if method == "MUS":
            start = rng.uniform(0, interval)
            cum, tot = [], 0.0
            for i in idx:
                tot += abs(rows[i]["amount"])
                cum.append(tot)
            for k in range(n):
                hp = start + k * interval
                j = min(bisect.bisect_left(cum, hp), len(idx) - 1)
                sample.append((idx[j], "MUS", round(hp, 2)))
            seen, uniq = set(), []
            for s in sample:                      # defensive: one line per item
                if s[0] not in seen:
                    seen.add(s[0])
                    uniq.append(s)
            sample = uniq
        elif method == "Random":
            for i in sorted(rng.sample(idx, n)):
                sample.append((i, "Random", None))
        elif method == "Systematic":
            step = len(idx) / n
            start = rng.uniform(0, step)
            for k in range(n):
                sample.append((idx[min(int(start + k * step), len(idx) - 1)], "Systematic", None))
    return dict(cats=cats, sample=sample, n=n, interval=interval, start=start, seed=seed, sign=sign,
                rmm=rmm, sap_eff=sap_eff, rf=rf, ef=ef, tm=tm, pm=pm, em=em, key_thr=key_thr,
                residual_value=R, residual_count=len(idx))


# ----------------------------------------------------------------------------- styling helpers
FONT = "Arial"
NAVY, WHITE, GREY, LGREY = "1F3864", "FFFFFF", "595959", "F2F2F2"
YELLOW, LBLUE, GREEN_F, RED_F = "FFFFCC", "DCE6F1", "C6EFCE", "FFC7CE"
BLUE, GREENTXT, REDTXT = "0000FF", "006100", "9C0006"
NUM = '#,##0;(#,##0);"-"'
NUM2 = '#,##0.00;(#,##0.00);"-"'
PCT = '0.0%'
thin = Side(style="thin", color="BFBFBF")


def put(ws, ref, v=None, fmt=None, bold=False, italic=False, color=None, fill=None, size=10, wrap=False,
        align=None, border=False):
    c = ws[ref]
    if v is not None:
        c.value = v
    c.font = Font(name=FONT, size=size, bold=bold, italic=italic, color=color)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = PatternFill("solid", start_color=fill)
    if wrap or align:
        c.alignment = Alignment(wrap_text=wrap, vertical="top", horizontal=align)
    if border:
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    return c


def band(ws, r, text, c1="B", c2="D"):
    for col in range(ord(c1), ord(c2) + 1):
        put(ws, f"{chr(col)}{r}", None, bold=True, color=WHITE, fill=NAVY)
    put(ws, f"{c1}{r}", text, bold=True, color=WHITE, fill=NAVY)


def inp(ws, ref, v, fmt=None, wrap=False):
    return put(ws, ref, v, fmt=fmt, color=BLUE, fill=YELLOW, wrap=wrap, border=True)


def head_row(ws, r, labels, start=1):
    for j, t in enumerate(labels):
        c = ws.cell(r, start + j, t)
        c.font = Font(name=FONT, size=10, bold=True, color=WHITE)
        c.fill = PatternFill("solid", start_color=NAVY)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[r].height = 30


def top_block(ws, title, last_col="H"):
    put(ws, "A1", title, bold=True, size=14, color=NAVY)
    for r, lab, f, fmt in ((2, "Client", "=Plan!$C$6", None), (3, "Period end", "=Plan!$C$7", "dd-mmm-yyyy"),
                           (4, "Audit area", "=Plan!$C$8", None)):
        put(ws, f"A{r}", lab, bold=True)
        put(ws, f"B{r}", f, fmt=fmt, color="008000", align="left")
    put(ws, "D2", "Tolerable misstatement", bold=True)
    put(ws, "F2", "=Plan!$C$19", fmt=NUM, color="008000")
    put(ws, "D3", "Method", bold=True)
    put(ws, "F3", "=Plan!$C$57", color="008000")


def landscape(ws):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


def status_cf(ws, rng, first):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({first},2)="OK"'], fill=PatternFill("solid", bgColor=GREEN_F), font=Font(name=FONT, color=GREENTXT, bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({first},5)="CHECK"'], fill=PatternFill("solid", bgColor=RED_F), font=Font(name=FONT, color=REDTXT, bold=True)))


# ----------------------------------------------------------------------------- workbook
def build(rows, P, out_path):
    res = select(rows, P)
    N = len(rows)
    cats = res["cats"]
    sign = res["sign"]
    method = P["method"]
    key_items = [i for i, c in enumerate(cats) if c.startswith("Key") or c.startswith("Top")]
    sample = res["sample"]
    n_slots = len(sample) if method != "Haphazard" else res["n"]

    P1, PL = 8, 7 + N                                # population rows
    K1, KL = 8, max(8, 7 + len(key_items))           # key item rows
    S1, SL = 8, max(8, 7 + n_slots)                  # sample rows
    pop = lambda col: f"Population!${col}${P1}:${col}${PL}"
    ki = lambda col: f"'Key Items'!${col}${K1}:${col}${KL}"
    sm = lambda col: f"Sample!${col}${S1}:${col}${SL}"

    wb = Workbook()
    wsP = wb.active
    wsP.title = "Plan"
    wsR = wb.create_sheet("Results")
    wsK = wb.create_sheet("Key Items")
    wsS = wb.create_sheet("Sample")
    wsPop = wb.create_sheet("Population")
    wsPar = wb.create_sheet("Parameters")

    # ---------------- Parameters
    ws = wsPar
    put(ws, "B2", "PARAMETERS - sampling tables", bold=True, size=14, color=NAVY)
    put(ws, "B3", "Defaults are ILLUSTRATIVE (Poisson-based reliability factors). Replace with your firm's methodology tables - "
                  "the Plan tab recalculates. Blue cells are editable.", italic=True, color=GREY)
    band(ws, 4, "A. Risk of material misstatement (RMM) - inherent risk vs test of controls", "B", "E")
    put(ws, "B5", "Inherent risk \\ Controls", bold=True, fill=LBLUE)
    for j, t in enumerate(TOC_OPTS):
        put(ws, f"{'CDE'[j]}5", t, bold=True, fill=LBLUE, wrap=True, align="center")
    for i, ir in enumerate(IR_OPTS):
        put(ws, f"B{6 + i}", ir, bold=True)
        for j in range(3):
            inp(ws, f"{'CDE'[j]}{6 + i}", RMM_MATRIX[ir][j])
    band(ws, 10, "B. Reliability (confidence) factor - RMM vs substantive analytical procedures", "B", "E")
    put(ws, "B11", "RMM \\ Analytics", bold=True, fill=LBLUE)
    for j, t in enumerate(SAP_OPTS):
        put(ws, f"{'CDE'[j]}11", t, bold=True, fill=LBLUE, align="center")
    for i, r_ in enumerate(RMM_ORDER):
        put(ws, f"B{12 + i}", r_, bold=True)
        for j in range(3):
            inp(ws, f"{'CDE'[j]}{12 + i}", RF_TABLE[r_][j], fmt="0.00")
    band(ws, 16, "C. Expansion factor (applied to expected misstatement)", "B", "E")
    put(ws, "B17", "RMM", bold=True, fill=LBLUE)
    put(ws, "C17", "Expansion factor", bold=True, fill=LBLUE)
    for i, r_ in enumerate(RMM_ORDER):
        put(ws, f"B{18 + i}", r_, bold=True)
        inp(ws, f"C{18 + i}", EF_TABLE[r_], fmt="0.0")
    band(ws, 22, "D. Conclusion rule for non-MUS methods (Random / Systematic / Haphazard)", "B", "E")
    put(ws, "B23", "Supports if known + projected misstatement <= this % of tolerable misstatement")
    inp(ws, "C23", SUPPORT_PCT, fmt="0%")
    notes = ["Notes",
             "1. Sample size = Residual population value x Reliability factor / (Tolerable misstatement - Expected misstatement x Expansion factor), rounded up.",
             "2. Reliability factor = -ln(risk of incorrect acceptance); e.g. 3.00 is about 95% confidence with zero expected misstatement.",
             "3. A significant risk forces RMM = High and gives no credit for substantive analytics.",
             "4. Controls reliance only reduces RMM when controls were tested AND found effective.",
             "5. MUS evaluation uses the Stringer bound (overstatements). Understatements are listed separately and must be evaluated by the auditor.",
             "6. Selection is made by the tool when run. If you change risk inputs in Excel, the required sample size updates but the selected items do not - re-run the tool."]
    for i, t in enumerate(notes):
        put(ws, f"B{25 + i}", t, bold=(i == 0), color=GREY if i else None)
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 34
    for c in "CDE":
        ws.column_dimensions[c].width = 26

    # ---------------- Plan
    ws = wsP
    put(ws, "B2", "AUDIT SAMPLING PLAN", bold=True, size=16, color=NAVY)
    put(ws, "B3", "Yellow cells with blue text are inputs. Everything else is a live formula.", italic=True, color=GREY)
    for r in range(4, 4):
        pass

    def line(r, label, value=None, note=None, kind="calc", fmt=None, wrap=False):
        put(ws, f"B{r}", label, bold=(kind == "calc" and False))
        if kind == "in":
            inp(ws, f"C{r}", value, fmt=fmt, wrap=wrap)
        elif kind == "static":
            put(ws, f"C{r}", value, fmt=fmt, border=True, wrap=wrap)
        else:
            put(ws, f"C{r}", value, fmt=fmt, border=True, wrap=wrap, align="right" if fmt else None)
        if note:
            put(ws, f"D{r}", note, italic=True, color=GREY, wrap=True)

    band(ws, 5, "A. ENGAGEMENT")
    line(6, "Client", P["client"], kind="in")
    line(7, "Period end", P["period_end"], kind="in", fmt="dd-mmm-yyyy")
    line(8, "Audit area / account", P["area"], kind="in")
    line(9, "Assertion(s) tested", P["assertion"], kind="in")
    line(10, "Objective of the test", P["objective"], kind="in", wrap=True)
    line(11, "Prepared by", P.get("prepared_by", ""), kind="in")
    line(12, "Prepared on", dt.datetime.now(), kind="in", fmt="dd-mmm-yyyy")
    line(13, "Reviewed by / date", None, kind="in")

    band(ws, 15, "B. MATERIALITY")
    line(16, "Overall materiality", float(P["om"]), "Determined in your materiality workpaper", "in", NUM)
    line(17, "Performance materiality %", P["pm_pct"], None, "in", "0%")
    line(18, "Performance materiality", "=C16*C17", None, "calc", NUM)
    if P.get("tm") in (None, ""):
        line(19, "Tolerable misstatement (this test)", "=C18", "Defaults to performance materiality; overwrite if different", "in", NUM)
    else:
        line(19, "Tolerable misstatement (this test)", float(P["tm"]), "Overwritten input (default = performance materiality)", "in", NUM)
    line(20, "Clearly trivial %", P["trivial_pct"], None, "in", "0%")
    line(21, "Clearly trivial threshold", "=C16*C20", None, "calc", NUM)

    band(ws, 23, "C. RISK ASSESSMENT AND ASSURANCE")
    line(24, "Inherent risk", P["ir"], "Low / Moderate / High", "in")
    line(25, "Significant risk?", "Yes" if P["sig"] else "No", "Yes = RMM High and no credit for analytics", "in")
    line(26, "Test of controls - result", P["toc"], "Reliance only if tested AND effective", "in")
    line(27, "Test of controls - description / WP ref", P.get("toc_desc", ""), None, "in", wrap=True)
    line(28, "Substantive analytics - effectiveness", P["sap"], "None / Limited / Effective", "in")
    line(29, "Analytics performed - description / WP ref", P.get("sap_desc", ""), None, "in", wrap=True)
    line(30, "Expected misstatement (amount)", float(P.get("em") or 0), "From prior year / known errors; 0 if none expected", "in", NUM)
    line(31, "Assessed RMM", '=IF(C25="Yes","High",INDEX(Parameters!$C$6:$E$8,MATCH(C24,Parameters!$B$6:$B$8,0),MATCH(C26,Parameters!$C$5:$E$5,0)))', "From Parameters tab")
    line(32, "Analytics credit applied", '=IF(C25="Yes","None",C28)')
    line(33, "Reliability (confidence) factor", "=INDEX(Parameters!$C$12:$E$14,MATCH(C31,Parameters!$B$12:$B$14,0),MATCH(C32,Parameters!$C$11:$E$11,0))", None, "calc", "0.00")
    line(34, "Expansion factor", "=INDEX(Parameters!$C$18:$C$20,MATCH(C31,Parameters!$B$18:$B$20,0))", None, "calc", "0.0")
    line(35, "Implied confidence level", "=1-EXP(-C33)", "1 minus the risk of incorrect acceptance", "calc", PCT)

    band(ws, 37, "D. POPULATION")
    line(38, "Total number of items", f"=COUNTA({pop('B')})", None, "calc", NUM)
    line(39, "Total value (absolute)", f"=SUM({pop('E')})", None, "calc", NUM)
    line(40, "Net book value", f"=SUM({pop('D')})", "Reconcile to the TB / lead schedule", "calc", NUM)
    line(41, "Population sign (+1 debit / -1 credit)", "=IF(C40>=0,1,-1)", "Used to read overstatements as positive", "calc", "0")
    line(42, "Zero-value items (not sampled)", f'=COUNTIF({pop("D")},0)', None, "calc", NUM)
    line(43, "Opposite-sign items (count)", f'=IF(C41=1,COUNTIF({pop("D")},"<0"),COUNTIF({pop("D")},">0"))', "e.g. credit balances in a debit population", "calc", NUM)

    band(ws, 45, "E. KEY ITEMS AND RESIDUAL POPULATION")
    line(46, "Key item threshold (absolute value)", res["key_thr"], "Items at or above are tested 100%. Input used at run time.", "static", NUM)
    line(47, "Other key item rules applied", "Specified IDs: " + (P.get("key_ids") or "none") + " | Opposite-sign items as key: " + ("Yes" if P.get("key_neg") else "No"), None, "static", wrap=True)
    line(48, "Key items - number", f'=COUNTIF({pop("F")},"Key item*")', None, "calc", NUM)
    line(49, "Key items - value", f'=SUMIF({pop("F")},"Key item*",{pop("E")})', None, "calc", NUM)
    line(50, "Top stratum (MUS) - number", f'=COUNTIF({pop("F")},"Top stratum*")', "MUS only: items >= sampling interval", "calc", NUM)
    line(51, "Top stratum (MUS) - value", f'=SUMIF({pop("F")},"Top stratum*",{pop("E")})', None, "calc", NUM)
    line(52, "Residual population - number", f'=COUNTIF({pop("F")},"Residual")', None, "calc", NUM)
    line(53, "Residual population - value", f'=SUMIF({pop("F")},"Residual",{pop("E")})', "Population from which the sample is drawn", "calc", NUM)
    line(54, "Check: key + top + residual + zero = total", f'=IF(ABS(C49+C51+C53-C39)<0.01,"OK","CHECK")', None, "calc")

    band(ws, 56, "F. SAMPLING")
    line(57, "Sampling method", method, METHOD_NOTES[method], "static")
    line(58, "Random seed (for reproducibility)", res["seed"], "Re-running with the same seed and file gives the same selection", "static")
    line(59, "Required sample size (residual items)", f'=IF(C53=0,0,MIN(C52,ROUNDUP(C53*C33/(C19-C30*C34),0)))', "Residual value x RF / (TM - EM x EF)", "calc", NUM)
    line(60, "MUS sampling interval", '=IF(AND(C57="MUS",C59>0),C53/C59,"n/a")', None, "calc", NUM)
    line(61, "Items in sample (Sample tab)", f'=SUMPRODUCT(--({sm("B")}<>""))', None, "calc", NUM)
    line(62, "Sample selection check", '=IF(C61=C59,"OK","CHECK - sample differs from required size; re-run tool if inputs changed")', None, "calc", wrap=True)
    line(63, "Sample value (absolute)", f"=SUM({sm('E')})", None, "calc", NUM)
    line(64, "Sample coverage of residual value", "=IF(C53=0,0,C63/C53)", None, "calc", PCT)
    line(65, "Residual value not tested", "=C53-C63", None, "calc", NUM)
    line(66, "Total value covered (key + top + sample)", "=C49+C51+C63", None, "calc", NUM)
    line(67, "Total coverage of population value", "=IF(C39=0,0,C66/C39)", None, "calc", PCT)

    band(ws, 69, "G. CONCLUSION")
    line(70, "Result of the test", "=Results!C24", "Enter audited values on the Key Items and Sample tabs", "calc", wrap=True)
    status_cf(ws, "C54", "C54")
    status_cf(ws, "C62", "C62")
    ws.conditional_formatting.add("C70", FormulaRule(formula=['LEFT(C70,8)="SUPPORTS"'], fill=PatternFill("solid", bgColor=GREEN_F), font=Font(name=FONT, color=GREENTXT, bold=True)))
    ws.conditional_formatting.add("C70", FormulaRule(formula=['OR(LEFT(C70,4)="DOES",LEFT(C70,4)="INCO")'], fill=PatternFill("solid", bgColor=RED_F), font=Font(name=FONT, color=REDTXT, bold=True)))
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 44
    ws.column_dimensions["C"].width = 34
    ws.column_dimensions["D"].width = 62
    for r in (10, 27, 29, 47, 62, 70):
        ws.row_dimensions[r].height = 30
    dv = lambda lst: DataValidation(type="list", formula1='"' + ",".join(lst) + '"', allow_blank=False)
    for ref, lst in (("C24", IR_OPTS), ("C25", ["Yes", "No"]), ("C26", TOC_OPTS), ("C28", SAP_OPTS)):
        d = dv(lst)
        ws.add_data_validation(d)
        d.add(ref)
    ws.freeze_panes = "A5"
    landscape(ws)

    # ---------------- Population
    ws = wsPop
    top_block(ws, "POPULATION")
    put(ws, "A5", "Category is set when the tool runs (key items / top stratum / residual / zero). Re-run the tool to change it.", italic=True, color=GREY)
    head_row(ws, 7, ["Line", "Item ID", "Description", "Book value", "Absolute value", "Category", "Cumulative residual (MUS)", "Testing status"])
    for k, r in enumerate(rows):
        rr = P1 + k
        c = cats[k]
        ws.cell(rr, 1, k + 1)
        ws.cell(rr, 2, r["id"])
        ws.cell(rr, 3, r["desc"])
        ws.cell(rr, 4, r["amount"]).number_format = NUM
        ws.cell(rr, 5, f"=ABS(D{rr})").number_format = NUM
        ws.cell(rr, 6, c)
        prev = "0" if k == 0 else f"G{rr - 1}"
        ws.cell(rr, 7, f'={prev}+IF($F{rr}="Residual",E{rr},0)').number_format = NUM
        ws.cell(rr, 8, f'=IF($F{rr}="Residual",IF(COUNTIF({sm("B")},$B{rr})>0,"In sample","Not selected"),IF($F{rr}="Zero value","Not tested","100% tested"))')
        for col in range(1, 9):
            ws.cell(rr, col).font = Font(name=FONT, size=10, color=("000000" if col not in (4,) else BLUE))
    ws.cell(PL + 1, 3, "Total").font = Font(name=FONT, bold=True)
    for col in (4, 5):
        L = "DE"[col - 4]
        c = ws.cell(PL + 1, col, f"=SUM({L}{P1}:{L}{PL})")
        c.font = Font(name=FONT, bold=True)
        c.number_format = NUM
    for L, w in zip("ABCDEFGH", (7, 18, 40, 16, 16, 26, 22, 16)):
        ws.column_dimensions[L].width = w
    ws.freeze_panes = "A8"
    ws.auto_filter.ref = f"A7:H{PL}"
    ws.conditional_formatting.add(f"A{P1}:H{PL}", FormulaRule(formula=[f'LEFT($F{P1},3)="Key"'], fill=PatternFill("solid", bgColor="FCE4D6")))
    ws.conditional_formatting.add(f"A{P1}:H{PL}", FormulaRule(formula=[f'LEFT($F{P1},3)="Top"'], fill=PatternFill("solid", bgColor="FFF2CC")))
    ws.conditional_formatting.add(f"A{P1}:H{PL}", FormulaRule(formula=[f'$H{P1}="In sample"'], fill=PatternFill("solid", bgColor=GREEN_F)))

    # ---------------- Key Items
    ws = wsK
    top_block(ws, "KEY ITEMS AND TOP STRATUM - tested 100%")
    put(ws, "A5", "Enter the audited value for every item (blue column). Misstatement is shown as overstatement (+) / understatement (-).", italic=True, color=GREY)
    head_row(ws, 7, ["#", "Item ID", "Description", "Book value", "Why selected", "Audited value", "Misstatement (+ over / - under)", "Procedure / WP ref / comments"])
    for k, i in enumerate(key_items):
        rr = K1 + k
        ws.cell(rr, 1, k + 1)
        ws.cell(rr, 2, rows[i]["id"])
        ws.cell(rr, 3, rows[i]["desc"])
        ws.cell(rr, 4, rows[i]["amount"]).number_format = NUM
        ws.cell(rr, 5, cats[i])
        f = ws.cell(rr, 6)
        f.number_format = NUM
        f.font = Font(name=FONT, color=BLUE)
        f.fill = PatternFill("solid", start_color=YELLOW)
        ws.cell(rr, 7, f'=IF(F{rr}="","",Plan!$C$41*(D{rr}-F{rr}))').number_format = NUM
        ws.cell(rr, 8).fill = PatternFill("solid", start_color=YELLOW)
    if not key_items:
        put(ws, "B8", "No key items identified.", italic=True)
    for L, w in zip("ABCDEFGH", (6, 18, 40, 16, 26, 16, 20, 40)):
        ws.column_dimensions[L].width = w
    ws.freeze_panes = "A8"

    # ---------------- Sample
    ws = wsS
    top_block(ws, "SAMPLE - selected from the residual population")
    msg = {"MUS": "MUS selection: hit points are positions in the cumulative residual value on the Population tab.",
           "Random": "Simple random selection (seeded).",
           "Systematic": "Systematic selection: every nth item from a random start.",
           "Haphazard": "Haphazard: enter the Item IDs you select (blue cells) from the Population tab. Avoid conscious bias."}[method]
    put(ws, "A5", msg + "  Seed: " + str(res["seed"]), italic=True, color=GREY)
    head_row(ws, 7, ["#", "Item ID", "Description", "Book value", "Absolute value", "Selection basis", "MUS hit point",
                     "Audited value", "Misstatement (+ over / - under)", "Tainting (over only)", "Rank", "Stringer bound contribution", "Procedure / WP ref / comments"])
    for k in range(n_slots):
        rr = S1 + k
        ws.cell(rr, 1, k + 1)
        b = ws.cell(rr, 2, rows[sample[k][0]]["id"] if method != "Haphazard" else None)
        if method == "Haphazard":
            b.font = Font(name=FONT, color=BLUE)
            b.fill = PatternFill("solid", start_color=YELLOW)
        ws.cell(rr, 3, f'=IF($B{rr}="","",IFERROR(INDEX({pop("C")},MATCH($B{rr}&"",{pop("B")},0)),"ID NOT FOUND"))')
        ws.cell(rr, 4, f'=IF($B{rr}="","",IFERROR(INDEX({pop("D")},MATCH($B{rr}&"",{pop("B")},0)),0))').number_format = NUM
        ws.cell(rr, 5, f"=IF(ISNUMBER(D{rr}),ABS(D{rr}),0)").number_format = NUM
        ws.cell(rr, 6, sample[k][1] if method != "Haphazard" else "Haphazard")
        ws.cell(rr, 7, sample[k][2] if method != "Haphazard" else None).number_format = NUM2
        f = ws.cell(rr, 8)
        f.number_format = NUM
        f.font = Font(name=FONT, color=BLUE)
        f.fill = PatternFill("solid", start_color=YELLOW)
        ws.cell(rr, 9, f'=IF(OR($B{rr}="",H{rr}=""),"",Plan!$C$41*(D{rr}-H{rr}))').number_format = NUM
        ws.cell(rr, 10, f'=IF(I{rr}="","",IF(D{rr}=0,0,MAX(0,MIN(1,I{rr}/ABS(D{rr})))))').number_format = "0.0%"
        ws.cell(rr, 11, f'=IF(N(J{rr})<=0,"",COUNTIF($J${S1}:$J${SL},">"&J{rr})+COUNTIF($J${S1}:J{rr},J{rr}))')
        ws.cell(rr, 12, f'=IF(OR(K{rr}="",NOT(ISNUMBER(Plan!$C$60))),"",J{rr}*(CHIINV(EXP(-Plan!$C$33),2*(K{rr}+1))/2-CHIINV(EXP(-Plan!$C$33),2*K{rr})/2)*Plan!$C$60)').number_format = NUM2
        ws.cell(rr, 13).fill = PatternFill("solid", start_color=YELLOW)
    if n_slots == 0:
        put(ws, "B8", "No sample required (no residual population).", italic=True)
    for L, w in zip("ABCDEFGHIJKLM", (6, 18, 38, 15, 15, 14, 14, 15, 18, 12, 7, 17, 36)):
        ws.column_dimensions[L].width = w
    ws.freeze_panes = "A8"

    # ---------------- Results
    ws = wsR
    top_block(ws, "RESULTS AND EVALUATION")
    put(ws, "A5", "Overstatement = book value higher than audited value (for a credit population, credits larger than audited).", italic=True, color=GREY)

    def rline(r, label, f, fmt=NUM, note=None, bold=False):
        put(ws, f"A{r}", label, bold=bold)
        put(ws, f"C{r}", f, fmt=fmt, border=True, bold=bold, align="right" if fmt else None)
        if note:
            put(ws, f"D{r}", note, italic=True, color=GREY)

    band(ws, 6, "A. KEY ITEMS AND TOP STRATUM (100% tested)", "A", "D")
    rline(7, "Items still to be audited", f'=SUMPRODUCT(--({ki("B")}<>""),--({ki("F")}=""))')
    rline(8, "Known overstatements", f'=SUMIF({ki("G")},">0")')
    rline(9, "Known understatements", f'=-SUMIF({ki("G")},"<0")')
    band(ws, 11, "B. SAMPLE (residual population)", "A", "D")
    rline(12, "Items still to be audited", f'=SUMPRODUCT(--({sm("B")}<>""),--({sm("H")}=""))')
    rline(13, "Number of items with overstatement", f'=COUNTIF({sm("I")},">0")')
    rline(14, "Overstatements found in sample", f'=SUMIF({sm("I")},">0")')
    rline(15, "Understatements found in sample", f'=-SUMIF({sm("I")},"<0")', note="Not included in the MUS bound - evaluate separately")
    band(ws, 17, "C. EVALUATION", "A", "D")
    rline(18, "Projected misstatement - MUS (sum of tainting x interval)", f'=IF(Plan!$C$57="MUS",SUM({sm("J")})*Plan!$C$60,"n/a")')
    rline(19, "Basic precision (reliability factor x interval)", '=IF(Plan!$C$57="MUS",Plan!$C$33*Plan!$C$60,"n/a")')
    rline(20, "Upper misstatement limit on residual (Stringer bound)", f'=IF(Plan!$C$57="MUS",SUM({sm("L")})+Plan!$C$33*Plan!$C$60,"n/a")', note="Basic precision + projected + incremental allowance")
    rline(21, "Projected misstatement - ratio method (non-MUS)", f'=IF(Plan!$C$57="MUS","n/a",IF(Plan!$C$63=0,0,C14/Plan!$C$63*Plan!$C$53))', note="Sample overstatement / sample value x residual value")
    rline(22, "Total: known + projected / upper limit", '=C8+IF(Plan!$C$57="MUS",C20,C21)', bold=True)
    rline(23, "Tolerable misstatement", "=Plan!$C$19")
    put(ws, "A24", "Conclusion", bold=True)
    put(ws, "C24", '=IF(C7+C12>0,"INCOMPLETE - enter audited values for all selected items",'
                   'IF(Plan!$C$57="MUS",IF(C22<=C23,"SUPPORTS - upper limit within tolerable misstatement","DOES NOT SUPPORT - upper limit exceeds tolerable misstatement"),'
                   'IF(C22<=Parameters!$C$23*C23,"SUPPORTS - projected misstatement well within tolerable misstatement",'
                   'IF(C22<=C23,"INCONCLUSIVE - close to tolerable misstatement; extend sample / further procedures","DOES NOT SUPPORT - projected misstatement exceeds tolerable misstatement"))))',
        bold=True, border=True, wrap=True)
    ws.merge_cells("C24:D24")
    ws.row_dimensions[24].height = 32
    ws.conditional_formatting.add("C24", FormulaRule(formula=['LEFT(C24,8)="SUPPORTS"'], fill=PatternFill("solid", bgColor=GREEN_F), font=Font(name=FONT, color=GREENTXT, bold=True)))
    ws.conditional_formatting.add("C24", FormulaRule(formula=['OR(LEFT(C24,4)="DOES",LEFT(C24,4)="INCO")'], fill=PatternFill("solid", bgColor=RED_F), font=Font(name=FONT, color=REDTXT, bold=True)))
    band(ws, 26, "D. FOR THE SUMMARY OF UNADJUSTED DIFFERENCES", "A", "D")
    rline(27, "Factual overstatements above clearly trivial", f'=SUMIF({ki("G")},">"&Plan!$C$21)+SUMIF({sm("I")},">"&Plan!$C$21)')
    rline(28, "Factual understatements above clearly trivial", f'=-SUMIF({ki("G")},"<"&-Plan!$C$21)-SUMIF({sm("I")},"<"&-Plan!$C$21)')
    rline(29, "Projected misstatement on residual (to be considered)", '=IF(Plan!$C$57="MUS",C18,C21)')
    put(ws, "A31", "If the result does not support the balance: consider extending the sample, performing further procedures on the population, "
                   "or proposing an adjustment. Document the basis in your file.", italic=True, color=GREY)
    ws.column_dimensions["A"].width = 56
    ws.column_dimensions["B"].width = 4
    ws.column_dimensions["C"].width = 26
    ws.column_dimensions["D"].width = 58
    for w_ in (wsK, wsS, wsPop):
        w_.sheet_properties.tabColor = "8EA9DB"
    wsP.sheet_properties.tabColor = NAVY
    wsR.sheet_properties.tabColor = "70AD47"
    wsPar.sheet_properties.tabColor = "A6A6A6"
    for w_ in (wsR, wsK, wsS, wsPop):
        landscape(w_)
    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)
    res["n_key"] = sum(1 for c in cats if c.startswith("Key"))
    res["n_top"] = sum(1 for c in cats if c.startswith("Top"))
    res["n_items"] = N
    return res


# ----------------------------------------------------------------------------- template
def make_template(path):
    rng = random.Random(7)
    wb = Workbook()
    ws = wb.active
    ws.title = "Instructions"
    lines = ["POPULATION TEMPLATE - Audit Sampling Tool", "",
             "1. Use the 'Population' tab: one row per item (invoice, voucher, journal line, asset ...).",
             "2. Columns: Item ID | Description | Book value.  Book value: debit positive, credit negative.",
             "3. Delete the sample rows and paste your population. Keep the header row.",
             "4. The population should reconcile to the trial balance / lead schedule for the account.",
             "5. Item IDs should be unique (duplicates are auto-numbered). CSV files with the same columns also work.",
             "", "The sample data is FICTITIOUS (trade receivables - open invoices)."]
    for i, t in enumerate(lines, 1):
        ws.cell(i, 1, t).font = Font(name=FONT, bold=(i == 1), size=12 if i == 1 else 10)
    ws.column_dimensions["A"].width = 100
    wp = wb.create_sheet("Population")
    for j, h in enumerate(["Item ID", "Description", "Book value"], 1):
        c = wp.cell(1, j, h)
        c.font = Font(name=FONT, bold=True, color=WHITE)
        c.fill = PatternFill("solid", start_color=NAVY)
    custs = ["Apex Traders", "Bharat Steel Ltd", "Capital Retail", "Delta Components", "Evergreen Foods", "Fortune Pharma",
             "Global Textiles", "Horizon Auto", "Indus Polymers", "Jupiter Electricals", "Kaveri Agro", "Lotus Paper"]
    for k in range(1, 151):
        amt = round(math.exp(rng.gauss(10.2, 1.1)), 0)
        if k in (12, 77):
            amt = round(amt * 12, 0)
        if k in (31, 108, 133):
            amt = -round(amt * 0.3, 0)
        wp.cell(k + 1, 1, f"INV-{k:04d}")
        wp.cell(k + 1, 2, f"{rng.choice(custs)}")
        c = wp.cell(k + 1, 3, amt)
        c.number_format = NUM
        c.font = Font(name=FONT, color=BLUE)
    wp.column_dimensions["A"].width = 14
    wp.column_dimensions["B"].width = 34
    wp.column_dimensions["C"].width = 16
    wb.save(path)
