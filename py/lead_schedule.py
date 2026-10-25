#!/usr/bin/env python3
"""
Lead Schedule Generator
=======================
Takes a MAPPED trial balance (each account already tagged with a standard FSLI
category and a sub-category) and builds an Excel audit file containing:

  Cover          client, year end, materiality, integrity checks, index
  Lead Summary   (1) lead by FS category  (2) lead by category + sub-category
  Balance Sheet  built by linking to the lead tabs
  Profit & Loss  built by linking to the lead tabs
  TB             the mapped trial balance (source data for every formula)
  A1, A2 ...     one lead tab per FS category: sub-category lead + account detail

Everything in the output is a live Excel formula - change a number on the TB tab
(or an adjustment) and every lead tab, the Balance Sheet and the P&L update.

Usage
-----
  python lead_schedule.py                      (interactive - asks you questions)
  python lead_schedule.py --make-template TB_Template.xlsx
  python lead_schedule.py --tb TB.xlsx --client "ABC Ltd" --year-end 2026-03-31 \
                          --materiality 5000000 [--currency INR] [--out Lead.xlsx]

TB input columns (header names are matched flexibly):
  Account Code | Account Name | FSLI Category | Sub-category | PY Balance |
  CY Balance | Adjustments (optional)
Sign convention: Debit = positive, Credit = negative.
"""
import argparse
import datetime as dt
import re
import sys
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover
    sys.exit("The 'openpyxl' package is missing.  Run:  python -m pip install openpyxl")

# --------------------------------------------------------------------------- #
# 1. STANDARD FSLI TAXONOMY  (edit here to change categories / order / names)
#    (statement, block, group label, ref letter, [(FSLI name, short tab name)])
# --------------------------------------------------------------------------- #
SECTIONS = [
    ("BS", "ASSETS", "Non-current assets", "A", [
        ("Property, Plant & Equipment", "PPE"),
        ("Intangible Assets", "Intangibles"),
        ("Non-current Investments", "NC Investments"),
        ("Other Non-current Assets", "Other NCA"),
    ]),
    ("BS", "ASSETS", "Current assets", "B", [
        ("Inventories", "Inventories"),
        ("Trade Receivables", "Receivables"),
        ("Cash & Bank", "Cash & Bank"),
        ("Current Investments", "Cur Investments"),
        ("Other Current Assets", "Other CA"),
    ]),
    ("BS", "EQUITY & LIABILITIES", "Equity", "C", [
        ("Share Capital", "Share Capital"),
        ("Reserves & Surplus", "Reserves"),
    ]),
    ("BS", "EQUITY & LIABILITIES", "Non-current liabilities", "D", [
        ("Long-term Borrowings", "LT Borrowings"),
        ("Deferred Tax Liabilities", "Deferred Tax"),
        ("Long-term Provisions", "LT Provisions"),
    ]),
    ("BS", "EQUITY & LIABILITIES", "Current liabilities", "E", [
        ("Short-term Borrowings", "ST Borrowings"),
        ("Trade Payables", "Payables"),
        ("Other Current Liabilities", "Other CL"),
        ("Short-term Provisions", "ST Provisions"),
    ]),
    ("PL", "PROFIT & LOSS", "Income", "F", [
        ("Revenue from Operations", "Revenue"),
        ("Other Income", "Other Income"),
    ]),
    ("PL", "PROFIT & LOSS", "Expenses", "G", [
        ("Cost of Materials / Purchases", "Cost of Sales"),
        ("Changes in Inventories", "Chg in Inventory"),
        ("Employee Benefit Expense", "Employee Cost"),
        ("Finance Costs", "Finance Costs"),
        ("Depreciation & Amortisation", "Depreciation"),
        ("Other Expenses", "Other Expenses"),
    ]),
    ("PL", "PROFIT & LOSS", "Tax", "H", [
        ("Tax Expense", "Tax"),
    ]),
]

# Presentation sign: +1 = show the TB sign (debit +), -1 = flip (credit shown +)
GROUP_SIGN = {"A": 1, "B": 1, "C": -1, "D": -1, "E": -1, "F": -1, "G": 1, "H": 1}

# Common alternative names accepted when reading a TB (case-insensitive)
ALIASES = {
    "ppe": "Property, Plant & Equipment", "fixed assets": "Property, Plant & Equipment",
    "property plant & equipment": "Property, Plant & Equipment",
    "intangibles": "Intangible Assets", "intangible": "Intangible Assets",
    "investments": "Non-current Investments", "long-term investments": "Non-current Investments",
    "inventory": "Inventories", "stock": "Inventories", "stocks": "Inventories",
    "debtors": "Trade Receivables", "sundry debtors": "Trade Receivables",
    "accounts receivable": "Trade Receivables", "receivables": "Trade Receivables",
    "cash & cash equivalents": "Cash & Bank", "cash & bank balances": "Cash & Bank",
    "cash": "Cash & Bank", "bank": "Cash & Bank",
    "equity share capital": "Share Capital", "capital": "Share Capital",
    "reserves": "Reserves & Surplus", "retained earnings": "Reserves & Surplus",
    "reserves & surplus": "Reserves & Surplus",
    "long term borrowings": "Long-term Borrowings", "term loans": "Long-term Borrowings",
    "short term borrowings": "Short-term Borrowings", "bank overdraft": "Short-term Borrowings",
    "creditors": "Trade Payables", "sundry creditors": "Trade Payables",
    "accounts payable": "Trade Payables", "payables": "Trade Payables",
    "other current liabilities": "Other Current Liabilities",
    "long term provisions": "Long-term Provisions", "short term provisions": "Short-term Provisions",
    "revenue": "Revenue from Operations", "sales": "Revenue from Operations",
    "turnover": "Revenue from Operations",
    "cost of goods sold": "Cost of Materials / Purchases", "cogs": "Cost of Materials / Purchases",
    "purchases": "Cost of Materials / Purchases", "cost of materials": "Cost of Materials / Purchases",
    "cost of materials consumed": "Cost of Materials / Purchases",
    "employee cost": "Employee Benefit Expense", "employee benefits": "Employee Benefit Expense",
    "finance cost": "Finance Costs", "interest expense": "Finance Costs",
    "depreciation": "Depreciation & Amortisation", "depreciation & amortization": "Depreciation & Amortisation",
    "income tax": "Tax Expense", "tax": "Tax Expense", "taxation": "Tax Expense",
}

# --------------------------------------------------------------------------- #
# 2. SMALL HELPERS
# --------------------------------------------------------------------------- #
FONT = "Arial"
NAVY, LIGHT, GREY, WHITE = "1F3864", "D9E2F3", "F2F2F2", "FFFFFF"
BLUE, GREEN, RED = "0000FF", "008000", "C00000"
NUM = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
DATEF = "dd-mmm-yyyy"


def norm(s):
    s = str(s or "").casefold().replace(" and ", " & ")
    return re.sub(r"\s+", " ", s).strip()


CANON = {norm(f): f for _, _, _, _, items in SECTIONS for f, _ in items}
ALIAS_N = {norm(k): v for k, v in ALIASES.items()}


def canon(raw):
    n = norm(raw)
    return CANON.get(n) or ALIAS_N.get(n)


# FSLI metadata: name -> dict
FSLI = {}
for _stmt, _block, _group, _letter, _items in SECTIONS:
    for _i, (_name, _short) in enumerate(_items, 1):
        _ref = f"{_letter}{_i}"
        FSLI[_name] = dict(stmt=_stmt, block=_block, group=_group, letter=_letter,
                           ref=_ref, tab=f"{_ref} {_short}"[:31])


def q(name):
    return "'" + name.replace("'", "''") + "'"


def put(ws, r, c, v=None, fmt=None, bold=False, italic=False, color=None, fill=None,
        h=None, wrap=False, size=10, top=False, bottom=False, indent=0):
    cell = ws.cell(row=r, column=c)
    if v is not None:
        cell.value = v
    cell.font = Font(name=FONT, size=size, bold=bold, italic=italic, color=color)
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = PatternFill("solid", start_color=fill, end_color=fill)
    if h or wrap or indent:
        cell.alignment = Alignment(horizontal=h, wrap_text=wrap, indent=indent, vertical="center")
    if top or bottom:
        cell.border = Border(top=Side(style="thin") if top else None,
                             bottom=Side(style="double" if bottom else "thin") if bottom else None)
    return cell


def band(ws, r, c1, c2, text, fill=NAVY, color=WHITE, bold=True):
    for c in range(c1, c2 + 1):
        put(ws, r, c, None, fill=fill, color=color, bold=bold)
    put(ws, r, c1, text, fill=fill, color=color, bold=bold)


def sumrefs(cells):
    return "+".join(cells) if cells else "0"


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def header_block(ws, title):
    put(ws, 1, 1, title, bold=True, size=14, color=NAVY)
    put(ws, 2, 1, "Client", bold=True)
    put(ws, 2, 2, "=Cover!$C$4", color=GREEN, h="left")
    put(ws, 3, 1, "Year ended", bold=True)
    put(ws, 3, 2, "=Cover!$C$5", fmt=DATEF, color=GREEN, h="left")
    put(ws, 4, 1, "Materiality", bold=True)
    put(ws, 4, 2, "=Cover!$C$8", fmt=NUM, color=GREEN, h="left")
    put(ws, 4, 3, "Performance mat.", bold=True, h="right")
    put(ws, 4, 4, "=Cover!$C$10", fmt=NUM, color=GREEN)
    put(ws, 4, 5, "Clearly trivial", bold=True, h="right")
    put(ws, 4, 6, "=Cover!$C$12", fmt=NUM, color=GREEN)


def col_headers(ws, r, labels, start=1):
    for i, lab in enumerate(labels):
        put(ws, r, start + i, lab, bold=True, color=WHITE, fill=NAVY,
            h="left" if i < 3 else "center", wrap=True)
    ws.row_dimensions[r].height = 30


def status_formats(ws, rng, first_cell):
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'{first_cell}="OK"'], font=Font(name=FONT, bold=True, color="006100"),
        fill=PatternFill("solid", start_color="C6EFCE", end_color="C6EFCE")))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'{first_cell}="CHECK"'], font=Font(name=FONT, bold=True, color="9C0006"),
        fill=PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE")))


# --------------------------------------------------------------------------- #
# 3. READ THE MAPPED TRIAL BALANCE
# --------------------------------------------------------------------------- #
HEADER_SYNONYMS = {
    "code": ["account code", "code", "a/c code", "ac code", "ledger code", "gl code", "account no", "account number"],
    "name": ["account name", "name", "ledger name", "account", "particulars", "description", "ledger"],
    "fsli": ["fsli", "fsli category", "category", "fs line item", "fs category", "fs line", "financial statement line item", "fs caption", "fsli name"],
    "sub": ["sub-category", "sub category", "subcategory", "sub-fsli", "sub fsli", "sub caption", "sub-caption"],
    "py": ["py balance", "prior year", "py", "previous year", "prior year balance", "py closing", "opening", "opening balance"],
    "cy": ["cy balance", "current year", "cy", "closing", "closing balance", "current year balance", "balance", "cy unadjusted"],
    "adj": ["adjustments", "adjustment", "adj", "audit adjustments", "adjustments (optional)"],
    "dr": ["debit", "dr", "cy debit"],
    "cr": ["credit", "cr", "cy credit"],
}


def num(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "").replace("₹", "").replace("$", "")
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    if s.upper().endswith("DR"):
        s = s[:-2].strip()
    elif s.upper().endswith("CR"):
        s = "-" + s[:-2].strip()
    try:
        x = float(s)
    except ValueError:
        raise ValueError(f"Cannot read '{v}' as a number")
    return -x if neg else x


def read_tb(path):
    wb = load_workbook(path, data_only=True)
    ws = wb["TB"] if "TB" in wb.sheetnames else wb.worksheets[0]
    hdr_row, cols = None, {}
    for r in range(1, 21):
        found = {}
        for c in range(1, ws.max_column + 1):
            val = norm(ws.cell(r, c).value)
            for role, syns in HEADER_SYNONYMS.items():
                if val in syns and role not in found.values():
                    found[c] = role
                    break
        roles = set(found.values())
        if "name" in roles and "fsli" in roles and ("cy" in roles or {"dr", "cr"} <= roles):
            hdr_row, cols = r, {role: c for c, role in found.items()}
            break
    if hdr_row is None:
        sys.exit("Could not find the header row in the TB. Expected columns like: Account Code, "
                 "Account Name, FSLI Category, Sub-category, PY Balance, CY Balance, Adjustments.")
    rows, skipped = [], 0
    for r in range(hdr_row + 1, ws.max_row + 1):
        g = lambda role: ws.cell(r, cols[role]).value if role in cols else None
        code, name, fsli = g("code"), g("name"), g("fsli")
        if code in (None, "") and name in (None, ""):
            continue
        if fsli in (None, "") and "total" in norm(name):
            skipped += 1
            continue
        try:
            cy = num(g("cy")) if "cy" in cols else num(g("dr")) - num(g("cr"))
            rows.append(dict(
                code=str(code).strip() if code not in (None, "") else "",
                name=str(name or "").strip(),
                fsli_raw=str(fsli or "").strip(),
                sub=str(g("sub") or "").strip() or "General",
                py=num(g("py")), cy=cy, adj=num(g("adj"))))
        except ValueError as e:
            sys.exit(f"TB row {r}: {e}")
    if not rows:
        sys.exit("No account rows found under the header row.")
    return rows, skipped


# --------------------------------------------------------------------------- #
# 4. BUILD THE WORKBOOK
# --------------------------------------------------------------------------- #
def build(tb, meta, out_path):
    wb = Workbook()
    ws_cv = wb.active
    ws_cv.title = "Cover"
    ws_ls = wb.create_sheet("Lead Summary")
    ws_bs = wb.create_sheet("Balance Sheet")
    ws_pl = wb.create_sheet("Profit & Loss")
    ws_tb = wb.create_sheet("TB")

    # ---- classify TB rows ------------------------------------------------- #
    accounts, subs, unmapped = {}, {}, []
    n = len(tb)
    T1, TL = 6, 5 + n          # first / last TB data row
    for i, a in enumerate(tb):
        a["row"] = T1 + i
        f = canon(a["fsli_raw"])
        a["fsli"] = f or a["fsli_raw"] or "(blank)"
        a["mapped"] = f is not None
        if f:
            accounts.setdefault(f, []).append(a)
            if a["sub"] not in subs.setdefault(f, []):
                subs[f].append(a["sub"])
        else:
            unmapped.append(a)

    present_groups = []     # (stmt, block, group, letter, [fsli names present])
    for stmt, block, group, letter, items in SECTIONS:
        names = [nm for nm, _ in items if nm in accounts]
        if names:
            present_groups.append((stmt, block, group, letter, names))

    def tbr(col):           # absolute TB range for a column letter
        return f"TB!${col}${T1}:${col}${TL}"

    def tab_total_row(f):   # row of the total in the FSLI lead tab (deterministic)
        return 9 + len(subs[f])

    # ======================= TB SHEET ======================================= #
    header_block(ws_tb, "Trial Balance (mapped) - source data")
    put(ws_tb, 3, 4, "Sign convention: debit = positive, credit = negative.  Blue = input, black = formula.",
        italic=True, color="595959")
    col_headers(ws_tb, 5, ["Account code", "Account name", "FSLI category", "Sub-category",
                           "PY balance", "CY balance (unadjusted)", "Adjustments",
                           "CY balance (final)", "Lead ref"])
    for a in tb:
        r = a["row"]
        put(ws_tb, r, 1, a["code"], color=BLUE)
        put(ws_tb, r, 2, a["name"], color=BLUE)
        put(ws_tb, r, 3, a["fsli"], color=BLUE if a["mapped"] else RED,
            fill=None if a["mapped"] else "FFC7CE")
        put(ws_tb, r, 4, a["sub"], color=BLUE)
        put(ws_tb, r, 5, a["py"], fmt=NUM, color=BLUE)
        put(ws_tb, r, 6, a["cy"], fmt=NUM, color=BLUE)
        put(ws_tb, r, 7, a["adj"], fmt=NUM, color=BLUE)
        put(ws_tb, r, 8, f"=F{r}+G{r}", fmt=NUM)
    tr = TL + 1
    put(ws_tb, tr, 2, "Total (should be nil)", bold=True, top=True)
    for c in range(5, 9):
        L = get_column_letter(c)
        put(ws_tb, tr, c, f"=SUM({L}{T1}:{L}{TL})", fmt=NUM, bold=True, top=True)
    widths(ws_tb, {"A": 14, "B": 38, "C": 30, "D": 26, "E": 16, "F": 18, "G": 15, "H": 18, "I": 10})
    ws_tb.freeze_panes = "A6"
    ws_tb.sheet_properties.tabColor = "7F7F7F"

    # ======================= LEAD SUMMARY =================================== #
    ws = ws_ls
    header_block(ws, "Lead Summary")
    LSH = ["Ref", "Financial statement category", "Lead tab", "PY balance", "CY unadjusted",
           "Adjustments", "CY final", "Variance (CY - PY)", "Variance %",
           "Variance > perf. mat.?", "Balance >= materiality?", "Tie-out diff"]
    band(ws, 6, 1, 12, "1.  LEAD SUMMARY BY FINANCIAL STATEMENT CATEGORY")
    col_headers(ws, 7, LSH)
    ls_row, group_sub, block_tot = {}, {}, {}
    r = 8
    cur_block = None
    block_groups = {}
    for stmt, block, group, letter, names in present_groups:
        if block != cur_block:
            if cur_block is not None:
                pass
            cur_block = block
            band(ws, r, 1, 12, block, fill=LIGHT, color=NAVY)
            r += 1
        put(ws, r, 2, group, bold=True, italic=True)
        r += 1
        first = r
        for f in names:
            info = FSLI[f]
            ls_row[f] = r
            put(ws, r, 1, info["ref"])
            put(ws, r, 2, f)
            put(ws, r, 3, f'=HYPERLINK("#{q(info["tab"])}!A1","{info["tab"]}")', color="0563C1")
            for col, tbc in ((4, "E"), (5, "F"), (6, "G")):
                put(ws, r, col, f"=SUMIFS({tbr(tbc)},{tbr('C')},$B{r})", fmt=NUM, color=GREEN)
            put(ws, r, 7, f"=E{r}+F{r}", fmt=NUM)
            put(ws, r, 8, f"=G{r}-D{r}", fmt=NUM)
            put(ws, r, 9, f'=IF(D{r}=0,"n/m",H{r}/ABS(D{r}))', fmt=PCT, h="right")
            put(ws, r, 10, f'=IF(Cover!$C$10=0,"",IF(ABS(H{r})>Cover!$C$10,"Explain",""))', h="center")
            put(ws, r, 11, f'=IF(Cover!$C$8=0,"",IF(ABS(G{r})>=Cover!$C$8,"Yes",""))', h="center")
            r += 1
        last = r - 1
        put(ws, r, 2, f"Total {group.lower()}", bold=True, top=True, fill=GREY)
        for col in range(4, 9):
            L = get_column_letter(col)
            put(ws, r, col, f"=SUM({L}{first}:{L}{last})", fmt=NUM, bold=True, top=True, fill=GREY)
        group_sub[letter] = r
        block_groups.setdefault(block, []).append(r)
        r += 2
    # block totals
    tot_labels = {"ASSETS": "TOTAL ASSETS", "EQUITY & LIABILITIES": "TOTAL EQUITY & LIABILITIES",
                  "PROFIT & LOSS": "NET LOSS / (PROFIT) FOR THE YEAR  (TB sign)"}
    summary_total_rows = []
    for block in ("ASSETS", "EQUITY & LIABILITIES", "PROFIT & LOSS"):
        put(ws, r, 2, tot_labels[block], bold=True, top=True, bottom=True, fill=LIGHT)
        for col in range(4, 9):
            L = get_column_letter(col)
            put(ws, r, col, "=" + sumrefs([f"{L}{x}" for x in block_groups.get(block, [])]),
                fmt=NUM, bold=True, top=True, bottom=True, fill=LIGHT)
        block_tot[block] = r
        summary_total_rows.append(r)
        r += 1
    # unmapped row
    um = r
    put(ws, um, 2, "Unmapped / unclassified balances", bold=True, color=RED)
    for col, tbc in ((4, "E"), (5, "F"), (6, "G")):
        L = get_column_letter(col)
        put(ws, um, col,
            f"=SUM({tbr(tbc)})-(" + sumrefs([f"{L}{x}" for blk in block_groups.values() for x in blk]) + ")",
            fmt=NUM, color=RED)
    put(ws, um, 7, f"=E{um}+F{um}", fmt=NUM, color=RED)
    put(ws, um, 8, f"=G{um}-D{um}", fmt=NUM, color=RED)
    r += 1
    chk = r
    put(ws, chk, 2, "Total of all categories (trial balance check - should be nil)", bold=True, top=True, bottom=True)
    for col in range(4, 8):
        L = get_column_letter(col)
        put(ws, chk, col, "=" + "+".join(f"{L}{x}" for x in summary_total_rows + [um]),
            fmt=NUM, bold=True, top=True, bottom=True)
    sec1_first, sec1_last = 8, chk
    r += 3

    # ---- section 2: by category and sub-category
    band(ws, r, 1, 12, "2.  LEAD BY FINANCIAL STATEMENT CATEGORY AND SUB-CATEGORY")
    r += 1
    col_headers(ws, r, ["Ref", "Financial statement category", "Sub-category", "PY balance", "CY unadjusted",
                        "Adjustments", "CY final", "Variance (CY - PY)", "Variance %",
                        "Variance > perf. mat.?", "Balance >= materiality?", "Tie-out diff"])
    r += 1
    sec2_first = r
    tie_cells = []
    for stmt, block, group, letter, names in present_groups:
        for f in names:
            info = FSLI[f]
            frow = r
            put(ws, r, 1, info["ref"], bold=True, fill=GREY)
            put(ws, r, 2, f, bold=True, fill=GREY)
            put(ws, r, 3, None, fill=GREY)
            sub_first = r + 1
            sub_last = r + len(subs[f])
            for col in range(4, 9):
                L = get_column_letter(col)
                if col == 7:
                    put(ws, r, col, f"=E{r}+F{r}", fmt=NUM, bold=True, fill=GREY)
                elif col == 8:
                    put(ws, r, col, f"=G{r}-D{r}", fmt=NUM, bold=True, fill=GREY)
                else:
                    put(ws, r, col, f"=SUM({L}{sub_first}:{L}{sub_last})", fmt=NUM, bold=True, fill=GREY)
            put(ws, r, 9, f'=IF(D{r}=0,"n/m",H{r}/ABS(D{r}))', fmt=PCT, h="right", bold=True, fill=GREY)
            put(ws, r, 10, None, fill=GREY)
            put(ws, r, 11, None, fill=GREY)
            s = ls_row[f]
            put(ws, r, 12, f"=ABS(D{r}-D{s})+ABS(E{r}-E{s})+ABS(F{r}-F{s})+ABS(G{r}-G{s})",
                fmt=NUM, fill=GREY)
            tie_cells.append(f"L{r}")
            r += 1
            for sname in subs[f]:
                put(ws, r, 3, sname, indent=1)
                for col, tbc in ((4, "E"), (5, "F"), (6, "G")):
                    put(ws, r, col, f"=SUMIFS({tbr(tbc)},{tbr('C')},$B${frow},{tbr('D')},$C{r})",
                        fmt=NUM, color=GREEN)
                put(ws, r, 7, f"=E{r}+F{r}", fmt=NUM)
                put(ws, r, 8, f"=G{r}-D{r}", fmt=NUM)
                put(ws, r, 9, f'=IF(D{r}=0,"n/m",H{r}/ABS(D{r}))', fmt=PCT, h="right")
                r += 1
    sec2_last = r - 1
    widths(ws, {"A": 8, "B": 44, "C": 26, "D": 15, "E": 15, "F": 14, "G": 15, "H": 17, "I": 11,
                "J": 14, "K": 15, "L": 12})
    ws.freeze_panes = "D8"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = NAVY
    # TB lead-ref lookup (needs LS rows)
    for a in tb:
        put(ws_tb, a["row"], 9,
            f"=IFERROR(INDEX('Lead Summary'!$A${sec1_first}:$A${sec1_last},"
            f"MATCH($C{a['row']},'Lead Summary'!$B${sec1_first}:$B${sec1_last},0)),\"UNMAPPED\")")

    # ======================= FSLI LEAD TABS ================================= #
    tab_diff_cells = []
    for stmt, block, group, letter, names in present_groups:
        for f in names:
            info = FSLI[f]
            w = wb.create_sheet(info["tab"])
            w.sheet_properties.tabColor = {"BS": "2F75B5", "PL": "548235"}[stmt] if letter in "ABCDEFGH" else "7F7F7F"
            header_block(w, f"Lead Schedule - {f}")
            put(w, 6, 1, info["ref"], bold=True, size=12, color=NAVY)
            put(w, 6, 2, f, bold=True, size=12, color=NAVY)
            put(w, 6, 5, '=HYPERLINK("#\'Lead Summary\'!A1","<< Back to Lead Summary")', color="0563C1")
            band(w, 7, 1, 11, "A.  Lead by sub-category")
            col_headers(w, 8, ["", "Sub-category", "", "PY balance", "CY unadjusted", "Adjustments",
                               "CY final", "Variance (CY - PY)", "Variance %", "WP ref", "Conclusion / comments"])
            r2 = 9
            for sname in subs[f]:
                put(w, r2, 2, sname)
                for col, tbc in ((4, "E"), (5, "F"), (6, "G")):
                    put(w, r2, col, f"=SUMIFS({tbr(tbc)},{tbr('C')},$B$6,{tbr('D')},$B{r2})",
                        fmt=NUM, color=GREEN)
                put(w, r2, 7, f"=E{r2}+F{r2}", fmt=NUM)
                put(w, r2, 8, f"=G{r2}-D{r2}", fmt=NUM)
                put(w, r2, 9, f'=IF(D{r2}=0,"n/m",H{r2}/ABS(D{r2}))', fmt=PCT, h="right")
                r2 += 1
            tot = r2
            assert tot == tab_total_row(f)
            put(w, tot, 2, f"Total {f}", bold=True, top=True, bottom=True, fill=LIGHT)
            for col in range(4, 9):
                L = get_column_letter(col)
                if col == 7:
                    put(w, tot, col, f"=E{tot}+F{tot}", fmt=NUM, bold=True, top=True, bottom=True, fill=LIGHT)
                elif col == 8:
                    put(w, tot, col, f"=G{tot}-D{tot}", fmt=NUM, bold=True, top=True, bottom=True, fill=LIGHT)
                else:
                    put(w, tot, col, f"=SUM({L}9:{L}{tot - 1})", fmt=NUM, bold=True, top=True, bottom=True, fill=LIGHT)
            put(w, tot, 9, f'=IF(D{tot}=0,"n/m",H{tot}/ABS(D{tot}))', fmt=PCT, h="right", bold=True,
                top=True, bottom=True, fill=LIGHT)
            s = ls_row[f]
            put(w, tot + 1, 2, "Per Lead Summary", italic=True)
            for col, L in ((4, "D"), (5, "E"), (6, "F"), (7, "G")):
                put(w, tot + 1, col, f"='Lead Summary'!{L}{s}", fmt=NUM, italic=True, color=GREEN)
            put(w, tot + 2, 2, "Difference (should be nil)", italic=True)
            put(w, tot + 2, 7, f"=ABS(D{tot}-D{tot + 1})+ABS(E{tot}-E{tot + 1})+ABS(F{tot}-F{tot + 1})+ABS(G{tot}-G{tot + 1})",
                fmt=NUM, italic=True)
            tab_diff_cells.append(f"{q(info['tab'])}!G{tot + 2}")
            # account detail
            d0 = tot + 4
            band(w, d0, 1, 11, "B.  Account detail (linked to the TB tab)")
            col_headers(w, d0 + 1, ["Account code", "Account name", "Sub-category", "PY balance",
                                    "CY unadjusted", "Adjustments", "CY final", "Variance (CY - PY)",
                                    "Variance %", "WP ref", "Tick / comments"])
            rr = d0 + 2
            for a in accounts[f]:
                t = a["row"]
                put(w, rr, 1, f"=TB!A{t}", color=GREEN)
                put(w, rr, 2, f"=TB!B{t}", color=GREEN)
                put(w, rr, 3, f"=TB!D{t}", color=GREEN)
                put(w, rr, 4, f"=TB!E{t}", fmt=NUM, color=GREEN)
                put(w, rr, 5, f"=TB!F{t}", fmt=NUM, color=GREEN)
                put(w, rr, 6, f"=TB!G{t}", fmt=NUM, color=GREEN)
                put(w, rr, 7, f"=E{rr}+F{rr}", fmt=NUM)
                put(w, rr, 8, f"=G{rr}-D{rr}", fmt=NUM)
                put(w, rr, 9, f'=IF(D{rr}=0,"n/m",H{rr}/ABS(D{rr}))', fmt=PCT, h="right")
                rr += 1
            dt_ = rr
            put(w, dt_, 2, "Total per account detail", bold=True, top=True, bottom=True, fill=LIGHT)
            for col in range(4, 9):
                L = get_column_letter(col)
                if col == 7:
                    put(w, dt_, col, f"=E{dt_}+F{dt_}", fmt=NUM, bold=True, top=True, bottom=True, fill=LIGHT)
                elif col == 8:
                    put(w, dt_, col, f"=G{dt_}-D{dt_}", fmt=NUM, bold=True, top=True, bottom=True, fill=LIGHT)
                else:
                    put(w, dt_, col, f"=SUM({L}{d0 + 2}:{L}{dt_ - 1})", fmt=NUM, bold=True, top=True, bottom=True, fill=LIGHT)
            put(w, dt_ + 1, 2, "Difference to lead total (should be nil)", italic=True)
            put(w, dt_ + 1, 7, f"=ABS(G{dt_}-G{tot})", fmt=NUM, italic=True)
            tab_diff_cells.append(f"{q(info['tab'])}!G{dt_ + 1}")
            widths(w, {"A": 14, "B": 40, "C": 24, "D": 15, "E": 15, "F": 14, "G": 15, "H": 17, "I": 11,
                       "J": 12, "K": 36})
            w.freeze_panes = "A9"
            w.sheet_view.showGridLines = False

    # ======================= BALANCE SHEET & P&L ============================ #
    def fs_sheet(ws, title, sign_note, blocks_spec):
        header_block(ws, title)
        put(ws, 5, 1, sign_note, italic=True, color="595959")
        col_headers(ws, 6, ["Ref", "Particulars", "", "CY (after adjustments)", "PY"])
        return 7

    def line(ws, r, f, sign):
        info = FSLI[f]
        t = tab_total_row(f)
        s = "-" if sign < 0 else ""
        put(ws, r, 1, f'=HYPERLINK("#{q(info["tab"])}!A1","{info["ref"]}")', color="0563C1")
        put(ws, r, 2, f, indent=1)
        put(ws, r, 4, f"={s}{q(info['tab'])}!G{t}", fmt=NUM, color=GREEN)
        put(ws, r, 5, f"={s}{q(info['tab'])}!D{t}", fmt=NUM, color=GREEN)

    def sub_total(ws, r, label, first, last, fill=GREY, big=False):
        put(ws, r, 2, label, bold=True, top=True, fill=fill, bottom=big)
        for col in (4, 5):
            L = get_column_letter(col)
            put(ws, r, col, f"=SUM({L}{first}:{L}{last})", fmt=NUM, bold=True, top=True, fill=fill, bottom=big)

    groups_by_letter = {g[3]: g for g in present_groups}

    # --- P&L first (BS links to it)
    pl = ws_pl
    r = fs_sheet(pl, "Profit & Loss Account", "Presented in natural signs (income and profit positive).  All figures link to the lead tabs.", None)
    pl_rows = {}
    for letter, label in (("F", "Income"), ("G", "Expenses")):
        g = groups_by_letter.get(letter)
        band(pl, r, 1, 5, label.upper(), fill=LIGHT, color=NAVY)
        r += 1
        first = r
        for f in (g[4] if g else []):
            line(pl, r, f, -1 if letter == "F" else 1)
            r += 1
        last = r - 1
        if g:
            sub_total(pl, r, f"Total {label.lower()}", first, last)
        else:
            put(pl, r, 2, f"Total {label.lower()}", bold=True, top=True, fill=GREY)
            put(pl, r, 4, 0, fmt=NUM, bold=True, top=True, fill=GREY)
            put(pl, r, 5, 0, fmt=NUM, bold=True, top=True, fill=GREY)
        pl_rows[letter] = r
        r += 2
    pbt = r
    put(pl, pbt, 2, "PROFIT / (LOSS) BEFORE TAX", bold=True, top=True, fill=LIGHT)
    for col in (4, 5):
        L = get_column_letter(col)
        put(pl, pbt, col, f"={L}{pl_rows['F']}-{L}{pl_rows['G']}", fmt=NUM, bold=True, top=True, fill=LIGHT)
    r += 2
    band(pl, r, 1, 5, "TAX", fill=LIGHT, color=NAVY)
    r += 1
    g = groups_by_letter.get("H")
    first = r
    for f in (g[4] if g else []):
        line(pl, r, f, 1)
        r += 1
    last = r - 1
    if g:
        sub_total(pl, r, "Total tax expense", first, last)
    else:
        put(pl, r, 2, "Total tax expense", bold=True, top=True, fill=GREY)
        put(pl, r, 4, 0, fmt=NUM, bold=True, top=True, fill=GREY)
        put(pl, r, 5, 0, fmt=NUM, bold=True, top=True, fill=GREY)
    taxrow = r
    r += 2
    pat = r
    put(pl, pat, 2, "PROFIT / (LOSS) FOR THE YEAR", bold=True, top=True, bottom=True, fill=LIGHT)
    for col in (4, 5):
        L = get_column_letter(col)
        put(pl, pat, col, f"={L}{pbt}-{L}{taxrow}", fmt=NUM, bold=True, top=True, bottom=True, fill=LIGHT)
    widths(pl, {"A": 8, "B": 50, "C": 4, "D": 22, "E": 18})
    pl.sheet_view.showGridLines = False
    pl.sheet_properties.tabColor = "548235"
    pl.freeze_panes = "A7"

    # --- Balance sheet
    bs = ws_bs
    r = fs_sheet(bs, "Balance Sheet", "Presented in natural signs (assets and liabilities positive).  All figures link to the lead tabs.", None)
    totals = {}
    for block, letters in (("ASSETS", ("A", "B")), ("EQUITY & LIABILITIES", ("C", "D", "E"))):
        band(bs, r, 1, 5, block)
        r += 1
        sub_rows = []
        for letter in letters:
            g = groups_by_letter.get(letter)
            grp_name = {"A": "Non-current assets", "B": "Current assets", "C": "Equity",
                        "D": "Non-current liabilities", "E": "Current liabilities"}[letter]
            put(bs, r, 2, grp_name, bold=True, italic=True)
            r += 1
            first = r
            for f in (g[4] if g else []):
                line(bs, r, f, GROUP_SIGN[letter])
                r += 1
            if letter == "C":     # current-year result sits in equity
                put(bs, r, 2, "Profit / (loss) for the year (from Profit & Loss)", indent=1)
                put(bs, r, 4, f"='Profit & Loss'!D{pat}", fmt=NUM, color=GREEN)
                put(bs, r, 5, f"='Profit & Loss'!E{pat}", fmt=NUM, color=GREEN)
                r += 1
            last = r - 1
            if last >= first:
                sub_total(bs, r, f"Total {grp_name.lower()}", first, last)
            else:
                put(bs, r, 2, f"Total {grp_name.lower()}", bold=True, top=True, fill=GREY)
                put(bs, r, 4, 0, fmt=NUM, bold=True, top=True, fill=GREY)
                put(bs, r, 5, 0, fmt=NUM, bold=True, top=True, fill=GREY)
            sub_rows.append(r)
            r += 2
        put(bs, r, 2, "TOTAL " + block, bold=True, top=True, bottom=True, fill=LIGHT)
        for col in (4, 5):
            L = get_column_letter(col)
            put(bs, r, col, "=" + "+".join(f"{L}{x}" for x in sub_rows), fmt=NUM, bold=True,
                top=True, bottom=True, fill=LIGHT)
        totals[block] = r
        r += 2
    bs_chk = r
    put(bs, bs_chk, 2, "Check: total assets less total equity & liabilities (should be nil)", italic=True)
    for col in (4, 5):
        L = get_column_letter(col)
        put(bs, bs_chk, col, f"={L}{totals['ASSETS']}-{L}{totals['EQUITY & LIABILITIES']}",
            fmt=NUM, italic=True)
    widths(bs, {"A": 8, "B": 56, "C": 4, "D": 22, "E": 18})
    bs.sheet_view.showGridLines = False
    bs.sheet_properties.tabColor = "2F75B5"
    bs.freeze_panes = "A7"

    # ======================= COVER ========================================== #
    cv = ws_cv
    put(cv, 2, 2, "AUDIT LEAD SCHEDULE PACK", bold=True, size=16, color=NAVY)
    put(cv, 3, 2, "Generated from the mapped trial balance - every figure is a live formula.", italic=True, color="595959")
    rows_in = [
        (4, "Client", meta["client"], None), (5, "Year ended", meta["year_end"], DATEF),
        (6, "Currency / units", meta["currency"], None), (7, "Materiality basis", meta["basis"], None),
        (8, "Overall materiality", meta["materiality"], NUM),
        (9, "Performance materiality %", meta["pm_pct"], "0%"),
    ]
    for rr_, lab, val, fmt in rows_in:
        put(cv, rr_, 2, lab, bold=True)
        put(cv, rr_, 3, val, fmt=fmt, color=BLUE, h="left", fill="FFFFE0")
    put(cv, 10, 2, "Performance materiality", bold=True)
    put(cv, 10, 3, "=C8*C9", fmt=NUM, h="left")
    put(cv, 11, 2, "Clearly trivial %", bold=True)
    put(cv, 11, 3, meta["trivial_pct"], fmt="0%", color=BLUE, h="left", fill="FFFFE0")
    put(cv, 12, 2, "Clearly trivial threshold", bold=True)
    put(cv, 12, 3, "=C8*C11", fmt=NUM, h="left")
    put(cv, 14, 2, "Prepared by", bold=True)
    put(cv, 14, 3, meta["prepared_by"], color=BLUE, h="left", fill="FFFFE0")
    put(cv, 15, 2, "Prepared on", bold=True)
    put(cv, 15, 3, meta["prepared_on"], fmt=DATEF, color=BLUE, h="left", fill="FFFFE0")
    put(cv, 16, 2, "Reviewed by", bold=True)
    put(cv, 16, 3, None, fill="FFFFE0")
    put(cv, 17, 2, "Reviewed on", bold=True)
    put(cv, 17, 3, None, fmt=DATEF, fill="FFFFE0", h="left")
    put(cv, 4, 5, "Yellow cells with blue text are inputs - change them here and every tab updates.",
        italic=True, color="595959")
    put(cv, 5, 5, "Materiality here is an INPUT you determined elsewhere (e.g. in your materiality workpaper).",
        italic=True, color="595959")

    band(cv, 19, 2, 4, "INTEGRITY CHECKS")
    put(cv, 20, 2, "Check", bold=True, fill=LIGHT)
    put(cv, 20, 3, "Difference", bold=True, fill=LIGHT, h="right")
    put(cv, 20, 4, "Status", bold=True, fill=LIGHT, h="center")
    checks = [
        ("Trial balance - CY unadjusted nets to nil", f"=TB!F{tr}"),
        ("Trial balance - PY nets to nil", f"=TB!E{tr}"),
        ("Adjustments net to nil", f"=TB!G{tr}"),
        ("Unmapped / unclassified balances (CY final)", f"='Lead Summary'!G{um}"),
        ("Lead tabs tie to Lead Summary (sum of differences)",
         "=" + sumrefs([f"'Lead Summary'!{c}" for c in tie_cells] + tab_diff_cells)),
        ("Balance sheet balances (CY)", f"='Balance Sheet'!D{bs_chk}"),
        ("Balance sheet balances (PY)", f"='Balance Sheet'!E{bs_chk}"),
    ]
    r = 21
    for lab, f in checks:
        put(cv, r, 2, lab)
        put(cv, r, 3, f, fmt=NUM, color=GREEN)
        put(cv, r, 4, f'=IF(ABS(C{r})<0.5,"OK","CHECK")', h="center")
        r += 1
    chk_first, chk_last = 21, r - 1
    status_formats(cv, f"D{chk_first}:D{chk_last}", f"D{chk_first}")
    put(cv, r, 2, "Overall", bold=True, top=True)
    put(cv, r, 4, f'=IF(COUNTIF(D{chk_first}:D{chk_last},"CHECK")=0,"OK","CHECK")', bold=True, h="center", top=True)
    status_formats(cv, f"D{r}", f"D{r}")
    r += 2
    band(cv, r, 2, 4, "CONTENTS")
    r += 1
    index = [("Lead Summary", "Lead by FS category, then by category and sub-category"),
             ("Balance Sheet", "Linked to the lead tabs"), ("Profit & Loss", "Linked to the lead tabs"),
             ("TB", "Mapped trial balance - source data for all formulas")]
    for stmt, block, group, letter, names in present_groups:
        for f in names:
            index.append((FSLI[f]["tab"], f))
    for tab, desc in index:
        put(cv, r, 2, f'=HYPERLINK("#{q(tab)}!A1","{tab}")', color="0563C1")
        put(cv, r, 3, desc)
        r += 1
    widths(cv, {"A": 3, "B": 46, "C": 24, "D": 12, "E": 70})
    cv.sheet_view.showGridLines = False
    cv.sheet_properties.tabColor = NAVY

    # ---- optional unmapped listing --------------------------------------- #
    if unmapped:
        w = wb.create_sheet("Unmapped")
        w.sheet_properties.tabColor = "C00000"
        put(w, 1, 1, "UNMAPPED ACCOUNTS - fix the FSLI category on the TB tab (see 'FSLI List' in the template)",
            bold=True, color=RED, size=12)
        col_headers(w, 3, ["Account code", "Account name", "FSLI text found", "CY balance (final)"])
        for i, a in enumerate(unmapped):
            put(w, 4 + i, 1, f"=TB!A{a['row']}", color=GREEN)
            put(w, 4 + i, 2, f"=TB!B{a['row']}", color=GREEN)
            put(w, 4 + i, 3, f"=TB!C{a['row']}", color=GREEN)
            put(w, 4 + i, 4, f"=TB!H{a['row']}", fmt=NUM, color=GREEN)
        widths(w, {"A": 14, "B": 40, "C": 34, "D": 20})

    # sheet order: Cover, Lead Summary, BS, P&L, TB, lead tabs..., Unmapped
    for ws_ in wb.worksheets:
        ws_.page_setup.orientation = "landscape"
        ws_.page_setup.fitToWidth = 1
        ws_.page_setup.fitToHeight = 0
        ws_.sheet_properties.pageSetUpPr.fitToPage = True
    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)
    return dict(unmapped=unmapped, n=n, tabs=sum(len(g[4]) for g in present_groups))


# --------------------------------------------------------------------------- #
# 5. TEMPLATE / SAMPLE TRIAL BALANCE
# --------------------------------------------------------------------------- #
def sample_tb():
    # (code, name, FSLI, sub-category, PY, CY, adjustment)   Dr +, Cr -
    S = [
        ("1001", "Freehold Land", "Property, Plant & Equipment", "Land", 42000, 42000, 0),
        ("1002", "Buildings (net)", "Property, Plant & Equipment", "Buildings", 71200, 68500, 0),
        ("1003", "Plant & Machinery (net)", "Property, Plant & Equipment", "Plant & Machinery", 88400, 95300, 0),
        ("1004", "Furniture & Fixtures (net)", "Property, Plant & Equipment", "Furniture & Equipment", 6900, 6200, 0),
        ("1005", "Capital Work-in-Progress", "Property, Plant & Equipment", "Capital WIP", 5100, 12400, 0),
        ("1010", "Computer Software (net)", "Intangible Assets", "Software", 4100, 3400, 0),
        ("1020", "Investment in Subsidiary", "Non-current Investments", "Unquoted equity", 15000, 15000, 0),
        ("1030", "Security Deposits", "Other Non-current Assets", "Deposits", 2100, 2100, 0),
        ("1100", "Raw Materials", "Inventories", "Raw materials", 34200, 38600, 0),
        ("1101", "Work in Progress", "Inventories", "Work in progress", 12800, 14900, 0),
        ("1102", "Finished Goods", "Inventories", "Finished goods", 36500, 41200, 0),
        ("1103", "Stores & Spares", "Inventories", "Stores & spares", 4800, 5300, 0),
        ("1200", "Trade Receivables - Domestic", "Trade Receivables", "Domestic", 55900, 62400, 0),
        ("1201", "Trade Receivables - Export", "Trade Receivables", "Export", 24600, 28700, 0),
        ("1202", "Provision for Doubtful Debts", "Trade Receivables", "Provision", -2600, -3200, -950),
        ("1300", "HDFC Bank - Current A/c", "Cash & Bank", "Bank balances", 14200, 18400, 0),
        ("1301", "SBI - Collection A/c", "Cash & Bank", "Bank balances", 5300, 6100, 0),
        ("1302", "Fixed Deposits", "Cash & Bank", "Fixed deposits", 15000, 20000, 0),
        ("1303", "Cash in Hand", "Cash & Bank", "Cash on hand", 120, 140, 0),
        ("1400", "Mutual Fund Investments", "Current Investments", "Mutual funds", 6000, 7500, 0),
        ("1500", "GST Input Credit", "Other Current Assets", "Statutory balances", 7800, 9300, 0),
        ("1501", "Advances to Suppliers", "Other Current Assets", "Advances", 4900, 5600, 0),
        ("1502", "Prepaid Expenses", "Other Current Assets", "Prepayments", 2300, 2700, 0),
        ("2001", "Equity Share Capital", "Share Capital", "Equity shares", -50000, -50000, 0),
        ("2002", "General Reserve", "Reserves & Surplus", "General reserve", -25000, -25000, 0),
        ("RE", "Retained Earnings - opening (balancing figure for demo)", "Reserves & Surplus", "Retained earnings", None, None, 0),
        ("2100", "Term Loan - HDFC (non-current)", "Long-term Borrowings", "Term loans", -72000, -60000, 0),
        ("2101", "Deferred Tax Liability", "Deferred Tax Liabilities", "Deferred tax", -8600, -9200, 0),
        ("2102", "Provision for Gratuity (non-current)", "Long-term Provisions", "Gratuity", -6500, -7400, 0),
        ("2200", "Working Capital Loan", "Short-term Borrowings", "Working capital", -30000, -34500, 0),
        ("2201", "Current Maturities of Term Loan", "Short-term Borrowings", "Current maturities", -12000, -12000, 0),
        ("2300", "Trade Payables - Goods", "Trade Payables", "Goods", -39200, -44800, 0),
        ("2301", "Trade Payables - Services", "Trade Payables", "Services", -8100, -9600, 0),
        ("2400", "Statutory Dues Payable", "Other Current Liabilities", "Statutory dues", -5400, -6200, 0),
        ("2401", "Employee Payables", "Other Current Liabilities", "Employee dues", -4700, -5100, 0),
        ("2402", "Advances from Customers", "Other Current Liabilities", "Customer advances", -2900, -3800, 0),
        ("2403", "Accrued Expenses", "Other Current Liabilities", "Accruals", -3900, -4500, -320),
        ("2500", "Provision for Gratuity (current)", "Short-term Provisions", "Gratuity", -1000, -1200, 0),
        ("2501", "Provision for Income Tax", "Short-term Provisions", "Income tax", -10400, -11600, 0),
        ("3001", "Sales - Domestic", "Revenue from Operations", "Domestic sales", -318000, -352000, 0),
        ("3002", "Sales - Export", "Revenue from Operations", "Export sales", -110300, -128500, 0),
        ("3003", "Scrap Sales", "Revenue from Operations", "Other operating revenue", -4100, -4700, 0),
        ("3100", "Interest Income", "Other Income", "Interest", -1700, -2100, 0),
        ("3101", "Gain on Mutual Funds", "Other Income", "Investment income", -480, -650, 0),
        ("3102", "Foreign Exchange Gain", "Other Income", "Forex", -900, -1300, 0),
        ("4001", "Raw Material Consumed", "Cost of Materials / Purchases", "Raw material consumed", 237000, 258400, 0),
        ("4002", "Freight & Duty on Purchases", "Cost of Materials / Purchases", "Freight & duty", 8900, 9800, 0),
        ("4003", "Change in Inventory - FG & WIP", "Changes in Inventories", "FG & WIP", -3100, -6800, 0),
        ("4100", "Salaries & Wages", "Employee Benefit Expense", "Salaries", 42100, 46300, 0),
        ("4101", "PF & ESI Contribution", "Employee Benefit Expense", "Statutory contributions", 4700, 5200, 0),
        ("4102", "Gratuity Expense", "Employee Benefit Expense", "Gratuity", 1800, 2100, 0),
        ("4200", "Interest - Term Loan", "Finance Costs", "Term loan interest", 8600, 7300, 0),
        ("4201", "Interest - Working Capital", "Finance Costs", "Working capital interest", 3200, 3800, 0),
        ("4202", "Bank Charges", "Finance Costs", "Bank charges", 380, 410, 0),
        ("4300", "Depreciation", "Depreciation & Amortisation", "Depreciation on PPE", 13100, 13900, 0),
        ("4301", "Amortisation of Software", "Depreciation & Amortisation", "Amortisation", 1300, 1400, 0),
        ("4400", "Power & Fuel", "Other Expenses", "Power & fuel", 12900, 14200, 0),
        ("4401", "Repairs & Maintenance", "Other Expenses", "Repairs & maintenance", 5300, 5900, 0),
        ("4402", "Freight Outward", "Other Expenses", "Selling & distribution", 9800, 11400, 0),
        ("4403", "Rent", "Other Expenses", "Rent", 3600, 3600, 0),
        ("4404", "Professional Fees", "Other Expenses", "Professional fees", 2400, 2900, 320),
        ("4405", "Bad Debts / ECL Provision", "Other Expenses", "Provisions & write-offs", 900, 1100, 950),
        ("4406", "Insurance", "Other Expenses", "Insurance", 1500, 1700, 0),
        ("4407", "Miscellaneous Expenses", "Other Expenses", "Miscellaneous", 3900, 4300, 0),
        ("4500", "Current Tax", "Tax Expense", "Current tax", 10400, 11600, 0),
        ("4501", "Deferred Tax", "Tax Expense", "Deferred tax", 450, 600, 0),
    ]
    py_sum = sum(x[4] for x in S if x[4] is not None)
    cy_sum = sum(x[5] for x in S if x[5] is not None)
    return [(c, n, f, s, (-py_sum if p is None else p), (-cy_sum if y is None else y), a)
            for c, n, f, s, p, y, a in S]


def make_template(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Instructions"
    lines = [
        ("TRIAL BALANCE TEMPLATE - Lead Schedule Generator", True),
        ("", False),
        ("1.  Go to the 'TB' tab. Delete the sample rows and paste your own mapped trial balance (keep the header row).", False),
        ("2.  Every account needs an 'FSLI category' - pick from the drop-down (list is on the 'FSLI List' tab).", False),
        ("3.  'Sub-category' is free text (e.g. Land, Buildings, Domestic, Export). Blank sub-categories become 'General'.", False),
        ("4.  Sign convention: DEBIT = positive, CREDIT = negative. Balances for each column should add up to nil.", False),
        ("5.  'PY balance' and 'Adjustments' are optional - leave as 0 if not available. Adjustments must net to nil.", False),
        ("6.  Save, then run the generator:   python lead_schedule.py   (or double-click Run_Lead_Schedule.bat)", False),
        ("", False),
        ("Header names are matched flexibly (e.g. 'Closing balance' is read as CY balance). If your TB has separate", False),
        ("Debit and Credit columns instead of one balance column, that works too (CY = Debit - Credit).", False),
        ("Accounts whose FSLI is not in the standard list are reported as UNMAPPED - nothing is silently dropped.", False),
        ("", False),
        ("The sample data on the TB tab is a fictitious manufacturing company - for illustration only.", False),
    ]
    for i, (t, b) in enumerate(lines, 1):
        put(ws, i, 1, t, bold=b, size=14 if b else 10, color=NAVY if b else None)
    ws.column_dimensions["A"].width = 120

    tb = wb.create_sheet("TB")
    col_headers(tb, 1, ["Account Code", "Account Name", "FSLI Category", "Sub-category",
                        "PY Balance", "CY Balance", "Adjustments"])
    for i, row in enumerate(sample_tb(), 2):
        for c, v in enumerate(row, 1):
            put(tb, i, c, v, fmt=NUM if c >= 5 else None, color=BLUE if c >= 5 else None)
    last = 1 + len(sample_tb())
    widths(tb, {"A": 14, "B": 52, "C": 30, "D": 26, "E": 14, "F": 14, "G": 14})
    tb.freeze_panes = "A2"
    lst = wb.create_sheet("FSLI List")
    col_headers(lst, 1, ["FSLI Category (use exactly)", "Statement", "Group", "Lead ref"])
    r = 2
    for stmt, block, group, letter, items in SECTIONS:
        for f, _ in items:
            put(lst, r, 1, f)
            put(lst, r, 2, "Balance Sheet" if stmt == "BS" else "Profit & Loss")
            put(lst, r, 3, group)
            put(lst, r, 4, FSLI[f]["ref"])
            r += 1
    widths(lst, {"A": 36, "B": 16, "C": 26, "D": 10})
    dv = DataValidation(type="list", formula1=f"='FSLI List'!$A$2:$A${r - 1}", allow_blank=True,
                        showErrorMessage=False)
    tb.add_data_validation(dv)
    dv.add(f"C2:C{max(last + 200, 500)}")
    wb.save(path)


# --------------------------------------------------------------------------- #
# 6. COMMAND LINE / INTERACTIVE FRONT END
# --------------------------------------------------------------------------- #
def parse_date(s):
    s = str(s).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d-%b-%Y", "%d %b %Y", "%d %B %Y", "%d-%m-%y", "%d/%m/%y"):
        try:
            return dt.datetime.strptime(s, fmt)
        except ValueError:
            pass
    raise ValueError(f"Could not read the date '{s}'. Use e.g. 31-03-2026 or 2026-03-31.")


def to_float(s):
    return float(str(s).replace(",", "").strip())


def run(tb_path, client, year_end, materiality, currency="", basis="", pm_pct=0.75,
        trivial_pct=0.05, prepared_by="", out=None):
    tb_path = Path(str(tb_path).strip().strip('"'))
    if not tb_path.exists():
        sys.exit(f"File not found: {tb_path}")
    rows, skipped = read_tb(tb_path)
    meta = dict(client=client, year_end=parse_date(year_end), currency=currency, basis=basis,
                materiality=to_float(materiality), pm_pct=pm_pct, trivial_pct=trivial_pct,
                prepared_by=prepared_by, prepared_on=dt.datetime.now())
    if out is None:
        safe = re.sub(r"[^\w\- ]+", "", client).strip().replace(" ", "_") or "Client"
        out = tb_path.with_name(f"{safe}_Lead_Schedules_{meta['year_end']:%Y%m%d}.xlsx")
    res = build(rows, meta, out)
    print(f"\nRead {res['n']} accounts" + (f" (skipped {skipped} total row[s])" if skipped else ""))
    print(f"Created {res['tabs']} lead tabs + Cover, Lead Summary, Balance Sheet, Profit & Loss, TB")
    if res["unmapped"]:
        print(f"\n!! {len(res['unmapped'])} account(s) have an FSLI category that is not in the standard list:")
        for a in res["unmapped"][:15]:
            print(f"     {a['code']:>8}  {a['name'][:40]:40}  FSLI text: '{a['fsli_raw']}'")
        print("   They are flagged red on the TB tab and listed on the 'Unmapped' tab.")
    print(f"\nSaved: {out}")
    print("Tip: open it in Excel - formulas calculate on open. Check the Cover tab's integrity checks first.")
    return out


def ask(prompt, default=None):
    s = input(f"{prompt}" + (f" [{default}]" if default not in (None, "") else "") + ": ").strip()
    return s or (default or "")


def interactive():
    print("=" * 62)
    print(" LEAD SCHEDULE GENERATOR")
    print("=" * 62)
    print("Tip: drag your mapped TB Excel file into this window to paste its path.\n")
    tb = ask("Mapped trial balance file (.xlsx)")
    client = ask("Client name")
    ye = ask("Year ended (e.g. 31-03-2026)")
    mat = ask("Overall materiality (number)")
    cur = ask("Currency / units (e.g. INR, INR '000)", "")
    basis = ask("Materiality basis (e.g. 5% of PBT)", "")
    pm = ask("Performance materiality % of overall", "75")
    tr = ask("Clearly trivial % of overall", "5")
    who = ask("Prepared by (initials)", "")
    try:
        run(tb, client, ye, mat, cur, basis, to_float(pm) / 100, to_float(tr) / 100, who)
    except ValueError as e:
        sys.exit(f"\nError: {e}")


def main():
    if len(sys.argv) == 1:
        interactive()
        return
    ap = argparse.ArgumentParser(description="Generate lead schedules from a mapped trial balance.")
    ap.add_argument("--make-template", metavar="FILE", help="write a sample TB template and exit")
    ap.add_argument("--tb")
    ap.add_argument("--client")
    ap.add_argument("--year-end")
    ap.add_argument("--materiality")
    ap.add_argument("--currency", default="")
    ap.add_argument("--basis", default="")
    ap.add_argument("--pm-pct", type=float, default=75.0, help="performance materiality as %% of overall (default 75)")
    ap.add_argument("--trivial-pct", type=float, default=5.0, help="clearly trivial as %% of overall (default 5)")
    ap.add_argument("--prepared-by", default="")
    ap.add_argument("--out")
    a = ap.parse_args()
    if a.make_template:
        make_template(a.make_template)
        print(f"Template written: {a.make_template}")
        return
    missing = [n for n in ("tb", "client", "year_end", "materiality") if not getattr(a, n)]
    if missing:
        ap.error("missing: " + ", ".join("--" + m.replace("_", "-") for m in missing))
    try:
        run(a.tb, a.client, a.year_end, a.materiality, a.currency, a.basis,
            a.pm_pct / 100, a.trivial_pct / 100, a.prepared_by, a.out)
    except ValueError as e:
        sys.exit(f"Error: {e}")


if __name__ == "__main__":
    main()
