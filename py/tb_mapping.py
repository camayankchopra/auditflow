#!/usr/bin/env python3
"""
TB Mapping
==========
Takes a RAW client trial balance (as exported from Tally / an ERP / Excel) and suggests an FSLI category and
sub-category for every account, so the file can go straight into the Lead Schedule tool.

How each account is mapped (in this order):
  1. Previous mapping   - last year's mapped TB / a mapping file, matched on account code, then on account name.
  2. Account name rules - ordered keyword rules (e.g. "Sundry Debtors" -> Trade Receivables), sign-aware.
  3. Ledger group       - if the TB has a Group / Under column (Tally style) it is used to confirm or overrule.
  4. Code range         - last resort, shown as LOW confidence.
Every suggestion carries a confidence (High / Medium / Low / Unmapped) and a reason. Nothing is silently guessed:
anything that is not High appears on the Review sheet.

Output workbook: Summary | TB (Lead Schedule format) | Review | Mapping (reusable) | FSLI List
"""
import csv
import datetime as dt
import io
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

import lead_schedule as ls
from lead_schedule import FSLI, SECTIONS, band, col_headers, put, widths

# --------------------------------------------------------------------------- #
# FSLI shorthand
# --------------------------------------------------------------------------- #
PPE, INT_, NCI, ONCA = "Property, Plant & Equipment", "Intangible Assets", "Non-current Investments", "Other Non-current Assets"
INV, TR, CB, CI, OCA = "Inventories", "Trade Receivables", "Cash & Bank", "Current Investments", "Other Current Assets"
SC, RS = "Share Capital", "Reserves & Surplus"
LTB, DTL, LTP = "Long-term Borrowings", "Deferred Tax Liabilities", "Long-term Provisions"
STB, TP, OCL, STP = "Short-term Borrowings", "Trade Payables", "Other Current Liabilities", "Short-term Provisions"
REV, OI = "Revenue from Operations", "Other Income"
COM, CHI, EMP = "Cost of Materials / Purchases", "Changes in Inventories", "Employee Benefit Expense"
FC, DEP, OE, TAX = "Finance Costs", "Depreciation & Amortisation", "Other Expenses", "Tax Expense"
UNMAPPED = None

GRADE = {"High": 3, "Medium": 2, "Low": 1, "Unmapped": 0}


def nm(s):
    """Normalise an account / group name for matching."""
    s = str(s or "").casefold().replace("&", " & ")
    s = re.sub(r"[^a-z0-9&%/ ]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


class Rule:
    __slots__ = ("rx", "fsli", "sub", "conf", "sign", "note", "alt")

    def __init__(self, pat, fsli, sub="", conf="High", sign=None, note="", alt=""):
        self.rx, self.fsli, self.sub, self.conf, self.sign, self.note, self.alt = re.compile(pat), fsli, sub, conf, sign, note, alt


R = Rule

# --------------------------------------------------------------------------- #
# ACCOUNT NAME RULES - first match wins, so specific rules come first.
# sign: "dr" = rule only applies to a debit balance, "cr" = credit balance.
# --------------------------------------------------------------------------- #
NAME_RULES = [
    # --- accounts that cannot be mapped by name
    R(r"suspense|clearing a/c|clearing account|unallocated|unidentified|inter ?(company|branch|unit)|reconciliation (difference|account)|control a/c|control account|difference in (books|tb)",
      UNMAPPED, note="Suspense / clearing / control account - cannot be mapped by name. Find out what it holds."),
    # --- statutory and tax assets
    R(r"\bprepaid|prepayments?|advance insurance|unexpired", OCA, "Prepayments"),
    R(r"rent receivable|accrued income|unbilled|contract assets?|insurance claim receivable|claims receivable", OCA, "Other receivables", sign="dr"),
    R(r"capital advances?|advances? for (capital|fixed assets?|plant|machinery|land)", ONCA, "Capital advances"),
    R(r"advances? tax|tax paid in advance|tds (receivable|recoverable|credit|deducted)|tcs (receivable|recoverable)|(gst|igst|cgst|sgst|vat) (input|receivable|recoverable|refund|credit)|input (tax )?(credit|gst|vat|igst|cgst|sgst)|\bitc\b|cenvat|income tax refund|refund (due|receivable)|mat credit|balances? with (government|statutory)|self assessment tax",
      OCA, "Statutory balances", sign="dr"),
    # --- advances received / accruals / other liabilities
    R(r"advances? (from|received from) customers?|customer advances?|customer deposits?|dealers? deposits?|distributors? deposits?|advances? received|unearned (revenue|income)|deferred (revenue|income)|income received in advance|deposits? (received|from (customers?|dealers?|distributors?|tenants?))|security deposits? (received|payable|from)",
      OCL, "Customer advances"),
    R(r"(salar(y|ies)|wages|bonus|remuneration|incentive|staff|employee|reimbursement|ex gratia) (payable|dues?|outstanding)|payable to (staff|employees)|employee payables?|accrued (salar|wage|bonus)|outstanding (salar|wage)|salar(y|ies) payable|wages payable|bonus payable",
      OCL, "Employee dues", sign="cr"),
    R(r"accrued (expenses|expense|liabilit)|outstanding (expenses|liabilit)|expenses payable|accruals?\b|provision for (expenses|expense|audit|bonus|services)|audit fees? payable|creditors? for (expenses|capital|services)|capital creditors?|payable for (expenses|capital)|other payables?|other (current )?liabilit",
      OCL, "Accruals"),
    R(r"(tds|tcs|gst|igst|cgst|sgst|vat|esi|esic|pf|epf|provident fund|professional tax|labou?r welfare fund|lwf|service tax|excise|customs duty|duty|sales tax) (payable|liability|output)|output (gst|igst|cgst|sgst|vat|tax)|gst (output|payable|liability)|statutory (dues|liabilit)|duties & taxes|taxes payable|withholding tax|dues? to government|tds on|tcs on|gst",
      OCL, "Statutory dues", sign="cr"),
    # --- tax
    R(r"deferred tax (liabilit|liabilities)|\bdtl\b", DTL, "Deferred tax"),
    R(r"deferred tax asset|\bdta\b", ONCA, "Deferred tax asset"),
    R(r"deferred tax", DTL, "Deferred tax", sign="cr"),
    R(r"deferred tax", ONCA, "Deferred tax asset", "Medium", sign="dr",
      note="Debit balance: deferred tax asset (other non-current assets). If this is the P&L charge, map to Tax Expense.", alt=TAX),
    R(r"(income|current|corporate|minimum alternate) tax (expense|charge|for the (year|period))|tax expense|taxation expense|\bcurrent tax\b|prior (year|period) tax|mat expense|income tax (earlier|prior)",
      TAX, "Current tax"),
    R(r"provision for (income|current|corporate)? ?tax(ation)?|income tax provision|tax provision|current tax liabilit|income tax payable",
      STP, "Income tax", sign="cr"),
    R(r"provision for (income|current|corporate)? ?tax(ation)?|income tax provision|tax provision",
      TAX, "Current tax", "Medium", sign="dr", note="Debit balance in a tax provision account: confirm whether it is the P&L charge.", alt=STP),
    R(r"\b(income|corporate) tax\b|\btaxation\b", TAX, "Current tax", "Medium", sign="dr",
      note="Generic tax account: confirm it is the P&L charge.", alt=STP),
    R(r"\b(income|corporate) tax\b|\btaxation\b", STP, "Income tax", "Medium", sign="cr",
      note="Generic tax account with a credit balance: confirm it is the liability.", alt=TAX),
    # --- provision / allowance for doubtful debts, bad debts
    R(r"provision for (doubtful|bad)|allowance for (doubtful|credit loss|bad)|expected credit loss|\becl\b|doubtful debts? provision|bad debts? provision",
      TR, "Provision", sign="cr"),
    R(r"provision for (doubtful|bad)|allowance for (doubtful|credit loss|bad)|expected credit loss|\becl\b",
      OE, "Bad debts", "Medium", sign="dr", note="Debit balance in a doubtful-debts provision: looks like the P&L charge (or a reversal). Confirm.", alt=TR),
    R(r"bad debts? (recovered|recovery)|recovery of bad", OI, "Bad debts recovered"),
    R(r"bad debts?|doubtful debts? expense|credit loss expense|debts? written off|irrecoverable", OE, "Bad debts"),
    # --- depreciation
    R(r"accumulated amorti[sz]ation|provision for amorti[sz]ation", INT_, "Accumulated amortisation"),
    R(r"accumulated depreciation|provision for depreciation|\bacc(um)?\.? ?dep|depreciation (reserve|fund|provision)",
      PPE, "Accumulated depreciation"),
    R(r"depreciation|amorti[sz]ation|impairment (loss|charge)|impairment of", DEP, "Depreciation"),
    # --- investment gains / losses
    R(r"(profit|gain) on (sale|redemption|disposal) of (mutual funds?|investments?|shares|securities|units|bonds)|fair value gain|(unrealised|unrealized|mtm|mark to market) gain",
      OI, "Gain on investments"),
    R(r"loss on (sale|redemption|disposal) of (mutual funds?|investments?|shares|securities|units|bonds)|fair value loss|(unrealised|unrealized|mtm|mark to market) loss|diminution in (the )?value of investments?",
      OE, "Loss on investments", "Medium", note="Loss on investments: confirm presentation (other expenses vs finance costs).", alt=FC),
    # --- income
    R(r"interest (accrued|payable|due)|accrued interest", OCL, "Interest accrued", sign="cr"),
    R(r"interest (accrued|receivable)|accrued interest|interest receivable", OCA, "Interest receivable", sign="dr"),
    R(r"export incentives?|duty drawback|incentive (received|income)|scrap sales?|sale of scrap|other operating (income|revenue)|job ?work (income|charges received)|royalty (income|received)",
      REV, "Other operating revenue", "Medium", note="Operating revenue or other income? Confirm per the entity's policy.", alt=OI),
    R(r"interest (income|received|earned|from)|interest on (fixed )?deposits?|interest on fd|interest on (bank|savings)|interest on loans? (given|advanced)|interest on (income tax )?refund|interest on investments?|dividend|rent(al)? (income|received|receipts?)|profit on sale of (fixed )?assets?|profit on (sale|disposal) of (property|vehicle|machinery)|gain on|foreign exchange gain|exchange (gain|difference gain)|discount received|commission (income|received)|brokerage received|sundry balances? written back|liabilit(y|ies) no longer|balances? written back|miscellaneous income|misc income|other income|insurance claim (received|income)|government grants?|grant income|subsidy|provision (no longer required|written back|reversed)|reversal of provision|write ?back|interest subsidy|rebate received",
      OI, "Other income"),
    R(r"loss on (sale|disposal|discard|write)|loss on", OE, "Loss on disposal"),
    R(r"sales? (promotion|commission|incentive|expenses?|travel|staff|office|tax payable)|selling (&|and) distribution|selling expenses?", OE, "Selling expenses"),
    R(r"\bsales?\b|revenue|turnover|service (income|charges|revenue|fees?)|(professional|consultancy|consulting) (fees? (received|income)|income)|fees? income|freight (income|recovered)|commission earned|installation (income|charges)|income from (operations|services|sale)|sale of (goods|products|services)|operating income|subscription (income|revenue)|training fees?|contract (revenue|income)",
      REV, "Sales"),
    # --- employee cost
    R(r"(?=.*\b(non current|long term)\b)(?=.*(provision for|gratuity|leave|compensated|pension|retirement|employee benefit))",
      LTP, "Employee benefits", sign="cr"),
    R(r"(?=.*\b(current|short term)\b)(?=.*(provision for|gratuity|leave|compensated|pension|retirement|employee benefit))",
      STP, "Employee benefits", sign="cr"),
    R(r"provision for (gratuity|leave|compensated|pension|retirement|employee benefits?)|gratuity (payable|provision|fund|liability)|leave (encashment|salary) (payable|provision|liability)|compensated absences|retirement benefits?|pension (liability|provision)",
      LTP, "Employee benefits", "Medium", sign="cr", note="Employee-benefit provision: confirm the current / non-current split.", alt=STP),
    R(r"provision for (warranty|litigation|contingenc|claims|onerous|expected losses|dividend|sales returns?)|proposed dividend|warranty provision",
      STP, "Provisions", "Medium", sign="cr", note="Provision: confirm whether it is current or non-current.", alt=LTP),
    R(r"dividend payable|unpaid dividend|unclaimed dividend", OCL, "Dividend payable"),
    R(r"contract(ual)? labou?r|casual labou?r|manpower", EMP, "Contract labour", "Medium",
      note="Contract labour: employee cost or other expenses depends on the entity's policy.", alt=OE),
    R(r"employees? (cost|benefits?|compensation|expenses?|welfare)|salar(y|ies)|wages|bonus|ex[- ]?gratia|staff (cost|welfare|salary|expenses|training|benefits?|amenities)|gratuity (expense|paid|contribution)|leave encashment (expense|paid)|(pf|provident fund|esic?) (employer|contribution|expense)|employer s? contribution|directors? s? remuneration|remuneration|incentive|overtime|labou?r (welfare|charges|cost)|stipend|leave travel|\blta\b|superannuation|workmen s? compensation",
      EMP, "Salaries & wages"),
    # --- finance costs
    R(r"interest (on|paid|expense|charges?|cost)|finance (cost|charges?)|borrowing costs?|loan processing|processing fees?|guarantee commission|\blc charges|forward (premium|contract)|\bncd\b|debenture interest|bill discounting charges|\binterest\b",
      FC, "Interest", sign="dr"),
    R(r"\binterest\b", OI, "Interest income", "Medium", sign="cr",
      note="Credit balance in an interest account: income or accrued liability? Confirm.", alt=OCL),
    R(r"bank charges?|bank commission|bank fees?|bank (service|processing) charges?|payment gateway (charges|fees)|\blc commission",
      OE, "Bank charges", "Medium", note="Bank charges: other expenses vs finance cost depends on the entity's policy.", alt=FC),
    # --- borrowings
    R(r"bank overdraft|\bbank od\b|\bod (a/c|account|limit)|overdraft|cash credit|\bcc (a/c|account|limit)|working capital (loan|facility|demand|limit)|\bwcdl\b|packing credit|\bpcfc\b|buyers? credit|bill discounting|short term (loan|borrowing|debt)|current (maturit|portion)|loans? repayable on demand|channel financ|invoice (discounting|financing)|factoring|credit card (payable|a/c)",
      STB, "Working capital", sign="cr"),
    R(r"term loan|long term (loan|borrowing|debt)|non current (loan|borrowing)|debentures?|\bbonds?\b|vehicle loan|car loan|home loan|housing loan|machinery loan|equipment loan|\becb\b|lease liabilit|hire purchase|\bncd\b",
      LTB, "Term loans", sign="cr"),
    R(r"secured loans?|unsecured loans?|\bloans?\b|borrowings?|directors? loan|promoters? loan|shareholders? loan|related part(y|ies) loan|inter ?corporate",
      LTB, "Loans", "Medium", sign="cr", note="Loan: confirm long-term vs short-term (repayable within 12 months is short-term).", alt=STB),
    R(r"advances? (to|paid to) (suppliers?|vendors?|creditors?)|supplier advances?|vendor advances?", OCA, "Advances to suppliers", sign="dr"),
    R(r"staff (loans?|advances?)|employees? (loans?|advances?)|salary advance|imprest|advances? to (employees?|staff)", OCA, "Staff advances", sign="dr"),
    R(r"loans? (and|&) advances|loans? (to|given|advanced)|advances? (to|given to) (directors?|related|subsidiar|others?)|advances? recoverable|inter ?corporate (loan|deposit) given|other advances|advances? paid|loan receivable|\badvances?\b",
      OCA, "Loans & advances", "Medium", sign="dr", note="Loans / advances: confirm current vs non-current and related-party disclosure.", alt=ONCA),
    # --- cost of materials, inventory movement
    R(r"(raw materials?|\brm\b|materials?) (consumed|consumption|used)|consumption of (raw )?materials?", COM, "Materials consumed"),
    R(r"(stores?|spares?|consumables?|packing materials?|packaging materials?|fuel|power) (consumed|consumption|used)|stores (and|&) spares consumed|consumption of stores",
      OE, "Stores & spares consumed"),
    R(r"opening (stock|inventor)|opening wip|decrease in (stock|inventor)|increase in (stock|inventor)|(increase|decrease|changes?) in (stock|inventor)|stock (adjustment|variance|write ?(off|down))|inventory (adjustment|write ?(off|down)|variance)|stock (written|loss)",
      CHI, "Changes in inventories"),
    R(r"purchases?|cost of (goods|sales|materials|revenue|production)|\bcogs\b|freight inward|carriage inward|inward (freight|carriage)",
      COM, "Purchases"),
    R(r"import (duty|clearing)|customs? duty|clearing (&|and) forwarding|direct (expenses|costs?|materials?)|manufacturing (expenses|costs?)|production (expenses|costs?)|processing charges|job ?work (charges|expenses|paid)|conversion charges|factory (expenses|overheads?)|material cost|sub ?contract",
      COM, "Direct costs", "Medium", note="Direct cost: cost of materials vs other expenses depends on the entity's policy.", alt=OE),
    # --- expense context (must beat the asset keywords below, e.g. "Repairs - Building")
    R(r"discount (allowed|given)|cash discount|trade discount|cd allowed", OE, "Discount allowed", "Medium",
      note="Discount allowed: expense or reduction of revenue per policy.", alt=REV),
    R(r"\brent\b(?! deposit)|rates (&|and) taxes|insurance|repairs?|maintenance|upkeep|travel|conveyance|telephone|telecom|internet|broadband|postage|courier|printing|stationery|office (expenses|supplies)|electricity (charges|expenses)|professional (fees?|charges)|legal|audit fees?|auditors?|consultan|advertis|marketing|promotion|commission paid|brokerage|freight|carriage|packing|donations?|\bcsr\b|subscriptions?|membership|licen[sc]e fees?|training|recruitment|security (charges|services)|housekeeping|cleaning|sundry expenses|general expenses|exchange (loss|fluctuation)|foreign exchange loss|loss on (sale|disposal|discard|write)|penalty|penalties|\bfines?\b|royalty|warranty|sitting fees?|entertainment|hospitality|vehicle (running|expenses|maintenance)|\bfuel\b|diesel|petrol|power (&|and) fuel|written off|write ?off|round(ing)? off|late fee|advisory|\bfees?\b|\bcharges?\b|\bexpenses?\b|\bexpense\b",
      OE, "Other expenses"),
    # --- capital WIP (before 'work in progress' inventory)
    R(r"capital work in progress|\bcwip\b|assets under (construction|development)|construction in progress|capital wip", PPE, "Capital WIP"),
    # --- inventory
    R(r"provision for (obsolete|slow|inventory|stock)|obsolescence|inventory (provision|reserve)|nrv provision", INV, "Provision"),
    R(r"inventor(y|ies)|stock[- ]in[- ]trade|stock in hand|closing stock|raw materials?|\brm\b|finished goods?|\bfg\b|work[- ]in[- ]progress|\bwip\b|stores( (&|and) spares)?\b|spares|packing materials?|packaging materials?|consumables?|traded goods|goods in transit|loose tools|semi finished|\bstocks?\b|merchandise|trading goods",
      INV, "Inventory"),
    # --- receivables, payables
    R(r"trade receivables?|sundry debtors?|\bdebtors?\b|accounts? receivable|receivables? (from|-)|customers?\b|bills? receivable|retention (money )?receivable|\ba/r\b",
      TR, "Trade debtors"),
    R(r"trade payables?|sundry creditors?|\bcreditors?\b|accounts? payable|payables? (to|for|-)|suppliers?\b|vendors?\b|bills? payable|\ba/p\b|msme (creditors?|payables?)|dues to (micro|msme|suppliers?)",
      TP, "Trade creditors"),
    # --- current investments (before bank names such as Axis / HDFC)
    R(r"mutual funds?|liquid funds?|money market|debt funds?|treasury bills?|\bt bills?\b|commercial papers?|certificates? of deposit|current investments?|short term investments?|quoted investments?|marketable securities",
      CI, "Mutual funds"),
    # --- cash and bank
    R(r"margin money|lien|held as security|under lien", CB, "Margin money / lien deposits", "Medium",
      note="Margin money / lien-marked deposits are often 'other bank balances' or non-current. Confirm.", alt=OCA),
    R(r"fixed deposits?|\bfds?\b|\bfdr\b|term deposits?|time deposits?|bank deposits?|recurring deposits?", CB, "Fixed deposits"),
    R(r"cash (in|on) hand|petty cash|cash balances?|cash (a/c|account|at (bank|office))|^cash$|cash (and|&) (bank|cash equivalents)|cash in (transit|safe|office)|cash float|imprest cash",
      CB, "Cash on hand"),
    R(r"bank|current (a/c|account|acc)|savings (a/c|account)|collection (a/c|account)|\bhdfc\b|\bicici\b|\bsbi\b|state bank|\baxis\b|kotak|yes bank|\bpnb\b|punjab national|canara|\bidbi\b|union bank|bank of (baroda|india|maharashtra)|indusind|\bidfc\b|federal bank|standard chartered|\bhsbc\b|citi ?bank|\brbl\b|bandhan|south indian bank|escrow|paytm|razorpay|payment gateway|wallet|paypal|stripe|cheques? (in|on) hand|remittance in transit|cash equivalents",
      CB, "Bank balances"),
    # --- investments
    R(r"subsidiar|associates?|joint ventures?|\bjv\b|investments? in (shares|equity|capital|partnership|llp|bonds|debentures|gold|property|properties|real estate)|unquoted (equity|investments?|shares)|non current investments?|long term investments?|investment property|\bnsc\b|\bppf\b|government securities|\bg ?sec\b|national savings|kisan vikas",
      NCI, "Unquoted equity"),
    R(r"investments?\b|shares (in|of|held)|debentures held|bonds held", NCI, "Investments", "Medium",
      note="Investment: confirm current vs non-current.", alt=CI),
    # --- fixed assets
    R(r"intangible assets under development|\biaud\b", INT_, "Under development"),
    R(r"goodwill|patents?|trade ?marks?|copyrights?|\bbrands?\b|intangibles?|computer software|software|website( development)?|technical know[- ]how|franchise|product development|development costs?|capitali[sz]ed development|customer (relationships|lists)|non compete|mining rights|licen[sc]es? (and|&) permits",
      INT_, "Software"),
    R(r"freehold land|leasehold land|\bland\b|\bbuildings?\b|factory (building|shed)|plant|machiner(y|ies)|furniture|fixtures?|vehicles?|motor (cars?|vehicles?|cycles?)|\bcars?\b|office equipments?|equipments?|computers?|laptops?|servers?|electrical (installations?|fittings?|equipments?)|air condition|leasehold improvements?|\btools?\b|\bdies\b|moulds?|tractors?|trucks?|generators?|\bdg sets?|\blifts?\b|elevators?|fire (fighting|equipment)|cctv|printers?|photocopiers?|renovation|interiors?|signage|solar|fixed assets?|tangible assets?|property plant|\bppe\b|right of use|\brou\b|gross block|net block|\bp ?& ?m\b|\bf ?& ?f\b",
      PPE, "Fixed assets"),
    # --- other assets
    R(r"security deposits?|rent deposits?|electricity deposits?|telephone deposits?|utility deposits?|lease deposits?|earnest money|\bemd\b", ONCA, "Deposits"),
    R(r"deposits?(?! received)", ONCA, "Deposits", "Medium", note="Deposit: confirm current vs non-current.", alt=OCA),
    R(r"preliminary expenses|pre ?operative|deferred (revenue )?expenditure|share issue expenses|misc(ellaneous)? expenditure", ONCA, "Deferred expenditure", "Medium",
      note="Unamortised expenditure shown as an asset: check it is allowed under the applicable framework.", alt=OE),
    R(r"other (current )?assets?|other receivables?|sundry (receivables?|assets?|balances?)|accrued income|unbilled|contract assets?|rent receivable|insurance claim receivable|claims receivable|due from (related|director|employees?)|amount recoverable|recoverable|receivable",
      OCA, "Other receivables", "Medium", sign="dr", note="Generic receivable: confirm the nature and current vs non-current.", alt=ONCA),
    # --- equity
    R(r"share premium|securities premium|premium on shares", RS, "Securities premium"),
    R(r"capital reserve|revaluation (reserve|surplus)|general reserve|reserves?|surplus|retained earnings|accumulated (profits?|losses|deficit)|profit (and|&) loss (a/c|account|appropriation)|p ?& ?l (a/c|account|appropriation)|opening balance equity|retained|undistributed|other comprehensive income|\boci\b|hedging reserve|foreign currency translation|debenture redemption reserve|capital redemption|statutory reserve|special reserve|share option|stock option outstanding",
      RS, "Reserves"),
    R(r"preference share|preferred (stock|shares)|\bcrps\b|\bocps?\b|optionally convertible|compulsorily convertible", SC, "Preference shares", "Medium",
      note="Preference shares: confirm equity vs liability classification.", alt=LTB),
    R(r"share application|share warrants?|shares? pending allotment|share suspense", SC, "Share application money", "Medium",
      note="Share application money pending allotment: check classification.", alt=OCL),
    R(r"share capital|capital stock|common stock|equity shares?|paid up capital|issued capital|\bshares\b", SC, "Equity shares"),
    R(r"partners? s? capital|proprietor|owner s? (capital|equity)|capital account|drawings?|partners? current|current account (of )?(partner|proprietor)|members? (capital|funds)|\bcapital\b|net worth|head office|branch (current )?account",
      SC, "Owner / partner capital", "Medium", note="Non-corporate capital: confirm presentation (capital vs reserves) and drawings.", alt=RS),
    # --- generic income
    R(r"\bincome\b|\bgains?\b|\breceipts?\b", OI, "Other income", "Medium", sign="cr", note="Generic income account: confirm operating revenue vs other income.", alt=REV),
]

# Sub-category refinement from the account name (first matching hint wins, otherwise the rule's default)
SUB_HINTS = {
    PPE: [(r"capital work|cwip|under construction", "Capital WIP"), (r"accumulated|provision for dep", "Accumulated depreciation"),
          (r"\bland\b", "Land"), (r"building|shed|renovation|interior", "Buildings"),
          (r"vehicle|\bcars?\b|motor|truck|tractor", "Vehicles"), (r"computer|laptop|server|printer", "Computers"),
          (r"furniture|fixture|office equipment|electrical|air cond|cctv|fire|\bf ?& ?f\b", "Furniture & Equipment"),
          (r"plant|machiner|generator|dg set|tool|\bdies\b|mould|lift|elevator|solar|\bp ?& ?m\b", "Plant & Machinery"),
          (r"leasehold", "Leasehold improvements"), (r"right of use|\brou\b", "Right-of-use assets")],
    INT_: [(r"software", "Software"), (r"goodwill", "Goodwill"), (r"patent|trade ?mark|copyright|brand", "IP & brands")],
    TR: [(r"export|foreign|overseas", "Export"), (r"domestic|local", "Domestic"), (r"related|group|subsidiar", "Related parties")],
    TP: [(r"service", "Services"), (r"goods|material", "Goods"), (r"related|group|subsidiar", "Related parties"), (r"msme|micro", "MSME")],
    INV: [(r"raw material|\brm\b", "Raw materials"), (r"finished|\bfg\b|traded|trading|merchandise", "Finished goods"),
          (r"work in progress|\bwip\b|semi", "Work in progress"), (r"stores|spares|consumable|packing|packaging|tools", "Stores & spares"),
          (r"provision|obsole", "Provision")],
    CB: [(r"fixed deposit|\bfds?\b|\bfdr\b|term deposit|time deposit|recurring", "Fixed deposits"), (r"cash", "Cash on hand")],
    RS: [(r"retained|accumulated|profit (and|&) loss|p ?& ?l|surplus", "Retained earnings"), (r"general reserve", "General reserve"),
         (r"premium", "Securities premium"), (r"revaluation", "Revaluation reserve")],
    LTB: [(r"term loan", "Term loans"), (r"debenture|bond|ncd", "Debentures / bonds"), (r"director|promoter|shareholder|related", "Related-party loans")],
    STB: [(r"overdraft|\bod\b", "Overdraft"), (r"cash credit|\bcc\b", "Cash credit"), (r"current (maturit|portion)", "Current maturities"),
          (r"working capital", "Working capital")],
    OCL: [(r"statutory|tds|gst|tcs|vat|pf|esi|duty|duties", "Statutory dues"), (r"employee|salary|wages|staff|bonus", "Employee dues"),
          (r"advance|unearned|deferred (rev|inc)", "Customer advances"), (r"accru|outstanding|expenses", "Accruals")],
    OCA: [(r"prepaid|prepayment", "Prepayments"), (r"gst|tds|tcs|vat|itc|tax|input", "Statutory balances"),
          (r"advance", "Advances"), (r"deposit", "Deposits")],
    OE: [(r"rent", "Rent"), (r"insurance", "Insurance"), (r"repair|maintenance", "Repairs & maintenance"),
         (r"travel|conveyance", "Travel"), (r"professional|legal|audit|consult", "Professional fees"),
         (r"advertis|marketing|promotion|selling", "Selling & marketing"), (r"bad debt|doubtful", "Bad debts"),
         (r"bank charges|bank commission", "Bank charges"), (r"telephone|internet|telecom|broadband", "Communication"),
         (r"electricity|power|fuel|water", "Power & fuel"), (r"freight|carriage|packing", "Freight & packing")],
    EMP: [(r"salar|wages|stipend", "Salaries & wages"), (r"bonus|incentive|ex gratia", "Bonus & incentive"),
          (r"gratuity|leave|pension|retirement|superannuation", "Retirement benefits"), (r"pf|provident|esi", "Contribution to funds"),
          (r"welfare", "Staff welfare")],
    REV: [(r"export", "Export sales"), (r"domestic|local", "Domestic sales"), (r"service|fees|consult", "Services"),
          (r"return", "Sales returns")],
}

# --------------------------------------------------------------------------- #
# LEDGER GROUP RULES (Tally "Under" column / ERP account group)
# --------------------------------------------------------------------------- #
GROUP_RULES = [
    (r"suspense", UNMAPPED, "Low"),
    (r"bank (od|occ|overdraft)|overdraft|cash credit", STB, "High"),
    (r"short term (loan|borrowing)", STB, "High"),
    (r"term loans?|long term (loan|borrowing)|debentures?", LTB, "High"),
    (r"secured loans?|unsecured loans?|loans? liability|borrowings?|^loans?$", LTB, "Medium"),
    (r"loans? (and|&) advances|loans? asset|advances? asset|deposits? asset|current assets", OCA, "Medium"),
    (r"sundry debtors?|trade receivables?|accounts? receivable|^debtors?$", TR, "High"),
    (r"sundry creditors?|trade payables?|accounts? payable|^creditors?$", TP, "High"),
    (r"bank accounts?|bank balances?|cash[- ]in[- ]hand|cash (and|&) (bank|cash equivalents)|cash equivalents|^cash$", CB, "High"),
    (r"capital work", PPE, "High"),
    (r"intangible", INT_, "High"),
    (r"fixed assets?|property plant|\bppe\b|tangible|plant (and|&) machinery|^land$|^buildings?$|furniture|^vehicles?$|^computers?$", PPE, "High"),
    (r"investments?", NCI, "Medium"),
    (r"stock[- ]in[- ]hand|inventor|closing stock|raw materials?|finished goods|^stores", INV, "High"),
    (r"sales accounts?|^revenue|turnover|income from operations", REV, "High"),
    (r"direct incomes?", REV, "Medium"),
    (r"indirect incomes?|other income|non operating income", OI, "High"),
    (r"employee|salaries|payroll", EMP, "High"),
    (r"finance cost|^interest|bank charges", FC, "High"),
    (r"depreciation", DEP, "High"),
    (r"purchase accounts?|^direct expenses?|^direct costs?|cost of (goods|materials|sales)|manufacturing|trading expenses", COM, "Medium"),
    (r"indirect expenses?|administrative|admin|selling|operating expenses|other expenses|expenses", OE, "Medium"),
    (r"duties (and|&) taxes|statutory (liabilit|dues)|taxes payable|^gst", OCL, "Medium"),
    (r"current liabilit", OCL, "Medium"),
    (r"provisions?", STP, "Medium"),
    (r"reserves?( (and|&) surplus)?|retained earnings|surplus", RS, "High"),
    (r"share capital|capital account|^capital|^equity", SC, "Medium"),
    (r"misc(ellaneous)? expenses? asset|deferred", ONCA, "Medium"),
]

CONTRA_RX = re.compile(r"return|accumulated|\bacc(um)?\b|provision|allowance|discount|rebate|drawing|reversal|written back|write ?back|credit note|contra|round|adjust")
EXPECT = {"A": 1, "B": 1, "G": 1, "H": 1, "C": -1, "D": -1, "E": -1, "F": -1}   # expected sign of the TB balance (Dr +)
SIGN_EXEMPT_FSLI = {RS, SC}                       # losses / drawings can legitimately flip these


def _sign_ok(sign, cy, py):
    v = cy if abs(cy) >= 0.5 else py
    if sign is None or abs(v) < 0.5:
        return True
    return v > 0 if sign == "dr" else v < 0


def pick_sub(fsli, name, default):
    if default in ("Other operating revenue", "Interest", "Interest income", "Interest accrued", "Interest receivable", "Gain on investments", "Loss on disposal"):
        return default
    for rx, sub in SUB_HINTS.get(fsli, []):
        if re.search(rx, name):
            return sub
    return default or "General"


def name_candidate(name, cy, py):
    for r in NAME_RULES:
        if r.sign and not _sign_ok(r.sign, cy, py):
            continue
        if r.rx.search(name):
            return r
    return None


def group_candidate(group):
    if not group:
        return None
    for pat, fsli, conf in GROUP_RULES:
        if re.search(pat, group):
            return fsli, conf
    return None


def code_fallback(code, cy, py):
    c = re.match(r"\d", str(code or "").strip())
    if not c:
        return None
    v = cy if abs(cy) >= 0.5 else py
    d = c.group(0)
    if d == "1" and v >= 0:
        return OCA
    if d == "2" and v <= 0:
        return OCL
    if d == "3" and v <= 0:
        return RS
    if d == "4" and v <= 0:
        return REV
    if d in "56789" and v >= 0:
        return OE
    return None


# --------------------------------------------------------------------------- #
# READING FILES
# --------------------------------------------------------------------------- #
EXTRA_SYN = {
    "code": ["ledger id", "acc code", "account id", "gl account", "a/c no", "sr no", "ledger no"],
    "name": ["ledger account", "account title", "account description", "gl description", "particular", "name of ledger", "name of account"],
    "group": ["group", "under", "parent", "primary group", "account group", "group name", "ledger group", "parent group", "account class", "class"],
    "py_dr": ["py debit", "py dr", "prior year debit", "previous year debit"],
    "py_cr": ["py credit", "py cr", "prior year credit", "previous year credit"],
    "drcr": ["dr/cr", "dr cr", "drcr", "balance type", "dr or cr", "type dr/cr"],
    "cy": ["closing bal", "net balance", "amount", "balance as at", "balance as on", "net", "tb balance", "closing balance (dr/cr)"],
}


def _syn(role):
    base = [ls.norm(x) for x in ls.HEADER_SYNONYMS.get(role, [])]
    return set(base + [ls.norm(x) for x in EXTRA_SYN.get(role, [])])


ROLE_SYN = {r: _syn(r) for r in ["code", "name", "fsli", "sub", "py", "cy", "adj", "dr", "cr", "group", "py_dr", "py_cr", "drcr"]}
ROLE_ORDER = ["py_dr", "py_cr", "drcr", "group", "fsli", "sub", "code", "name", "adj", "py", "cy", "dr", "cr"]


def _grid(data, filename):
    """Return the worksheet cells as a list of row lists (xlsx or csv)."""
    ext = Path(filename).suffix.lower()
    if ext == ".csv":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        return [list(r) for r in csv.reader(io.StringIO(text))]
    wb = load_workbook(io.BytesIO(data), data_only=True)
    pref = [n for n in wb.sheetnames if n.casefold() in ("tb", "trial balance", "mapping")]
    ws = wb[pref[0]] if pref else wb.worksheets[0]
    return [[c for c in row] for row in ws.iter_rows(values_only=True)]


def _find_header(grid, need):
    for ri, row in enumerate(grid[:30]):
        found = {}
        for ci, v in enumerate(row):
            val = ls.norm(v)
            if not val:
                continue
            for role in ROLE_ORDER:
                if val in ROLE_SYN[role] and role not in found.values():
                    found[ci] = role
                    break
        roles = set(found.values())
        if need(roles):
            return ri, {role: ci for ci, role in found.items()}
    return None, None


def read_raw_tb(data, filename, cr_positive=False):
    grid = _grid(data, filename)
    hi, cols = _find_header(grid, lambda r: "name" in r and ("cy" in r or {"dr", "cr"} <= r))
    if hi is None:
        raise ValueError("Could not find the header row. The trial balance needs an account name column and either a balance "
                         "column or separate Debit and Credit columns (for example: Account Code, Account Name, Debit, Credit).")
    layout = "balance column" if "cy" in cols else "debit / credit columns"
    rows = []
    for ri in range(hi + 1, len(grid)):
        row = grid[ri]
        g = lambda role: row[cols[role]] if role in cols and cols[role] < len(row) else None
        code, name = g("code"), g("name")
        if (code in (None, "")) and (name in (None, "")):
            continue
        sname = str(name or "").strip()
        if not sname:
            continue
        low = ls.norm(sname)
        if code in (None, "") and re.search(r"\b(grand )?total\b|^difference$|^net (profit|loss)$|^closing balance$", low):
            continue
        try:
            if "cy" in cols:
                cy = ls.num(g("cy"))
                ind = str(g("drcr") or "").strip().upper()
                if ind.startswith("C") and cy > 0:
                    cy = -cy
                elif ind.startswith("D") and cy < 0:
                    cy = -cy
                if cr_positive:
                    cy = -cy
            else:
                cy = ls.num(g("dr")) - ls.num(g("cr"))
            if "py_dr" in cols or "py_cr" in cols:
                py = ls.num(g("py_dr")) - ls.num(g("py_cr"))
            else:
                py = ls.num(g("py"))
                if cr_positive and "py" in cols:
                    py = -py
        except ValueError as e:
            raise ValueError(f"Row {ri + 1} ({sname}): {e}")
        rows.append(dict(code=str(code).strip() if code not in (None, "") else "", name=sname,
                         group=str(g("group") or "").strip(), cy=cy, py=py))
    if not rows:
        raise ValueError("No account rows found under the header row.")
    return rows, dict(layout=layout, has_group="group" in cols, has_py=("py" in cols or "py_dr" in cols))


def read_previous(data, filename):
    """Previous mapping: any sheet with (code or name) + FSLI columns. Returns (by_code, by_name, n_valid, n_invalid)."""
    grid = _grid(data, filename)
    hi, cols = _find_header(grid, lambda r: "fsli" in r and ("name" in r or "code" in r))
    if hi is None:
        raise ValueError("The previous mapping file needs an FSLI Category column and an Account Name (or Account Code) column.")
    by_code, by_name, ok, bad = {}, {}, 0, 0
    for row in grid[hi + 1:]:
        g = lambda role: row[cols[role]] if role in cols and cols[role] < len(row) else None
        f = ls.canon(g("fsli"))
        if g("fsli") in (None, ""):
            continue
        if not f:
            bad += 1
            continue
        sub = str(g("sub") or "").strip() or "General"
        code, name = str(g("code") or "").strip(), nm(g("name"))
        if code:
            by_code[code] = (f, sub)
        if name:
            by_name[name] = (f, sub)
        ok += 1
    return by_code, by_name, ok, bad


# --------------------------------------------------------------------------- #
# MAPPING ENGINE
# --------------------------------------------------------------------------- #
def map_accounts(rows, prev=None):
    by_code, by_name = (prev[0], prev[1]) if prev else ({}, {})
    for r in rows:
        name, group = nm(r["name"]), nm(r["group"])
        cy, py = r["cy"], r["py"]
        r.update(fsli=None, sub="", conf="Unmapped", basis="No match", note="", alt="")
        # 1. previous mapping
        hit = by_code.get(r["code"]) if r["code"] else None
        how = "Previous mapping (code)"
        if not hit and name in by_name:
            hit, how = by_name[name], "Previous mapping (name)"
        if hit:
            r.update(fsli=hit[0], sub=hit[1], conf="High", basis=how, note="Same as the previous mapping.")
        else:
            rule = name_candidate(name, cy, py)
            grp = group_candidate(group)
            if rule is not None and rule.fsli is UNMAPPED:
                r.update(note=rule.note, basis="Not mapped")
            elif rule and grp and grp[0] == rule.fsli:
                conf = "High" if (rule.conf == "High" or grp[1] == "High") else "Medium"
                r.update(fsli=rule.fsli, sub=pick_sub(rule.fsli, name, rule.sub), conf=conf, basis="Account name + ledger group",
                         note=rule.note if conf != "High" else "", alt=rule.alt)
            elif rule and grp and grp[0] is not UNMAPPED:
                clash = f"Name suggests '{rule.fsli}', ledger group '{r['group']}' suggests '{grp[0]}'. Confirm."
                if (rule.conf == "High" and grp[1] == "Medium") or (rule.fsli == INT_ and grp[0] == PPE):  # generic group never overrules a specific name
                    r.update(fsli=rule.fsli, sub=pick_sub(rule.fsli, name, rule.sub), conf="High", basis="Account name",
                             note="", alt=rule.alt)
                elif rule.conf == "Medium" and grp[1] == "High":
                    r.update(fsli=grp[0], sub=pick_sub(grp[0], name, ""), conf="Medium", basis="Name and ledger group disagree",
                             note=clash, alt=rule.fsli)
                else:
                    r.update(fsli=rule.fsli, sub=pick_sub(rule.fsli, name, rule.sub), conf="Medium",
                             basis="Name and ledger group disagree", note=clash, alt=grp[0])
            elif rule:
                r.update(fsli=rule.fsli, sub=pick_sub(rule.fsli, name, rule.sub), conf=rule.conf, basis="Account name",
                         note=rule.note, alt=rule.alt)
            elif grp and grp[0] is not UNMAPPED:
                r.update(fsli=grp[0], sub=pick_sub(grp[0], name, ""), conf=grp[1], basis="Ledger group",
                         note="" if grp[1] == "High" else f"Mapped from the ledger group '{r['group']}' only.")
            else:
                fb = code_fallback(r["code"], cy, py)
                if fb:
                    r.update(fsli=fb, sub="General", conf="Low", basis="Code range / balance direction only",
                             note="No name or group match. Guess based on the account code and balance direction. Confirm.")
                else:
                    r.update(note="No rule matched this account name. Choose the FSLI category yourself.")
        # sign sanity (only when something is mapped)
        if r["fsli"]:
            exp = EXPECT.get(FSLI[r["fsli"]]["letter"])
            v = cy if abs(cy) >= 0.5 else py
            if exp and abs(v) >= 0.5 and (v > 0) != (exp > 0) and r["fsli"] not in SIGN_EXEMPT_FSLI \
                    and not CONTRA_RX.search(name):
                side = "debit" if v > 0 else "credit"
                msg = f"Balance is a {side}, which is unusual for {r['fsli']}. Check the mapping."
                r["note"] = (r["note"] + " " + msg).strip()
                if r["conf"] == "High":
                    r["conf"] = "Medium"
    return rows


# --------------------------------------------------------------------------- #
# OUTPUT WORKBOOK
# --------------------------------------------------------------------------- #
FILL = {"Medium": "FFF2CC", "Low": "FCE4D6", "Unmapped": "FFC7CE"}
NUMF = ls.NUM


def build(rows, meta, out_path, info=None, prev_stats=None):
    info = info or {}
    n = len(rows)
    wb = Workbook()
    sm = wb.active
    sm.title = "Summary"
    tb = wb.create_sheet("TB")
    rv = wb.create_sheet("Review")
    mp = wb.create_sheet("Mapping")
    lst = wb.create_sheet("FSLI List")

    # ---- FSLI list
    col_headers(lst, 1, ["FSLI Category (use exactly)", "Statement", "Group", "Lead ref"])
    r = 2
    for stmt, block, group, letter, items in SECTIONS:
        for f, _ in items:
            put(lst, r, 1, f)
            put(lst, r, 2, "Balance Sheet" if stmt == "BS" else "Profit & Loss")
            put(lst, r, 3, group)
            put(lst, r, 4, FSLI[f]["ref"])
            r += 1
    fsli_last = r - 1
    widths(lst, {"A": 36, "B": 16, "C": 26, "D": 10})

    # ---- TB (Lead Schedule format)
    hdr = ["Account Code", "Account Name", "FSLI Category", "Sub-category", "PY Balance", "CY Balance", "Adjustments",
           "Map confidence", "Map basis", "Review note"]
    col_headers(tb, 1, hdr)
    for i, x in enumerate(rows, 2):
        put(tb, i, 1, x["code"] or None)
        put(tb, i, 2, x["name"])
        put(tb, i, 3, x["fsli"], fill=FILL.get(x["conf"]))
        put(tb, i, 4, x["sub"] or None)
        put(tb, i, 5, x["py"], fmt=NUMF, color=ls.BLUE)
        put(tb, i, 6, x["cy"], fmt=NUMF, color=ls.BLUE)
        put(tb, i, 7, 0, fmt=NUMF, color=ls.BLUE)
        put(tb, i, 8, x["conf"], fill=FILL.get(x["conf"]), h="center")
        put(tb, i, 9, x["basis"])
        put(tb, i, 10, (x["note"] + (f"  Alternative: {x['alt']}." if x["alt"] else "")).strip() or None, wrap=False)
    last = n + 1
    tr = last + 1
    put(tb, tr, 2, "Total (should be nil)", bold=True)
    for c in (5, 6, 7):
        L = "EFG"[c - 5]
        put(tb, tr, c, f"=SUM({L}2:{L}{last})", fmt=NUMF, bold=True, top=True)
    widths(tb, {"A": 14, "B": 46, "C": 30, "D": 24, "E": 14, "F": 14, "G": 14, "H": 14, "I": 32, "J": 80})
    tb.freeze_panes = "C2"
    tb.auto_filter.ref = f"A1:J{last}"
    dv = DataValidation(type="list", formula1=f"='FSLI List'!$A$2:$A${fsli_last}", allow_blank=True, showErrorMessage=False)
    tb.add_data_validation(dv)
    dv.add(f"C2:C{last + 200}")

    # ---- Review
    todo = [(i, x) for i, x in enumerate(rows, 2) if x["conf"] != "High"]
    todo.sort(key=lambda t: -abs(t[1]["cy"] if abs(t[1]["cy"]) >= 0.5 else t[1]["py"]))
    put(rv, 1, 1, "Accounts to review", bold=True, size=14, color=ls.NAVY)
    put(rv, 2, 1, "Everything not rated High is listed, largest balance first. Fix the FSLI on the TB sheet; this list follows it.", italic=True, color=ls.GREY)
    col_headers(rv, 4, ["#", "Account code", "Account name", "CY balance", "Mapped to (TB sheet)", "Confidence", "Why / what to check",
                        "Alternative to consider", "Your decision"])
    for k, (ri, x) in enumerate(todo, 1):
        rr = 4 + k
        put(rv, rr, 1, k)
        put(rv, rr, 2, f"=IF(TB!A{ri}=\"\",\"\",TB!A{ri})", color=ls.GREEN)
        put(rv, rr, 3, f"=TB!B{ri}", color=ls.GREEN)
        put(rv, rr, 4, f"=TB!F{ri}", fmt=NUMF, color=ls.GREEN)
        put(rv, rr, 5, f"=IF(TB!C{ri}=\"\",\"(not mapped)\",TB!C{ri})", color=ls.GREEN)
        put(rv, rr, 6, x["conf"], fill=FILL.get(x["conf"]), h="center")
        put(rv, rr, 7, x["note"] or x["basis"], wrap=True)
        put(rv, rr, 8, x["alt"] or None, wrap=True)
        put(rv, rr, 9, None, fill="FFFF99")
    if not todo:
        put(rv, 5, 1, "Nothing to review - every account was mapped with High confidence.", bold=True, color=ls.GREEN)
    widths(rv, {"A": 5, "B": 14, "C": 40, "D": 14, "E": 30, "F": 12, "G": 70, "H": 30, "I": 28})
    rv.freeze_panes = "A5"

    # ---- Mapping (reusable next year)
    put(mp, 1, 1, "Reusable mapping - upload this file as 'Previous mapping' next year or for the same client", bold=True, color=ls.NAVY)
    put(mp, 2, 1, "Snapshot of the suggestions at the time of generation. If you changed the TB sheet, upload the TB sheet instead.", italic=True, color=ls.GREY)
    col_headers(mp, 4, ["Account Code", "Account Name", "FSLI Category", "Sub-category"])
    for k, x in enumerate(rows, 5):
        put(mp, k, 1, x["code"] or None)
        put(mp, k, 2, x["name"])
        put(mp, k, 3, x["fsli"])
        put(mp, k, 4, x["sub"] or None)
    widths(mp, {"A": 14, "B": 50, "C": 32, "D": 26})
    mp.freeze_panes = "A5"

    # ---- Summary
    put(sm, 1, 2, "TB Mapping - Summary", bold=True, size=14, color=ls.NAVY)
    put(sm, 2, 2, "Client", bold=True)
    put(sm, 2, 3, meta.get("client") or "Client")
    put(sm, 3, 2, "Year ended", bold=True)
    if meta.get("year_end"):
        put(sm, 3, 3, meta["year_end"], fmt=ls.DATEF, h="left")
    put(sm, 4, 2, "Input layout", bold=True)
    put(sm, 4, 3, info.get("layout", "") + (", with ledger group column" if info.get("has_group") else ""))
    put(sm, 5, 2, "Previous mapping", bold=True)
    put(sm, 5, 3, f"{prev_stats[0]} accounts loaded" + (f", {prev_stats[1]} rows had an unrecognised FSLI" if prev_stats[1] else "") if prev_stats else "Not used")

    band(sm, 7, 2, 6, "Mapping results")
    rng = lambda col: f"TB!${col}$2:${col}${last}"
    items = [("Accounts read", f"=COUNTA({rng('B')})"),
             ("High confidence", f"=COUNTIF({rng('H')},\"High\")"),
             ("Medium confidence (review)", f"=COUNTIF({rng('H')},\"Medium\")"),
             ("Low confidence (review)", f"=COUNTIF({rng('H')},\"Low\")"),
             ("Unmapped (choose a category)", f"=COUNTIF({rng('H')},\"Unmapped\")")]
    for k, (lab, f) in enumerate(items, 8):
        put(sm, k, 2, lab)
        put(sm, k, 3, f, fmt="#,##0", h="right")
    put(sm, 14, 2, "Mapped from previous mapping")
    put(sm, 14, 3, f"=COUNTIF({rng('I')},\"Previous mapping*\")", fmt="#,##0", h="right")
    put(sm, 15, 2, "Mapped from account name / ledger group")
    put(sm, 15, 3, f"=COUNTIF({rng('I')},\"Account name*\")+COUNTIF({rng('I')},\"Ledger group*\")+COUNTIF({rng('I')},\"Name and ledger group*\")", fmt="#,##0", h="right")

    band(sm, 17, 2, 6, "Integrity checks")
    checks = [("Trial balance CY nets to nil", f"=TB!F{tr}"), ("Trial balance PY nets to nil", f"=TB!E{tr}"),
              ("Accounts with no FSLI category", f"=COUNTBLANK(TB!$C$2:$C${last})")]
    for k, (lab, f) in enumerate(checks, 18):
        put(sm, k, 2, lab)
        put(sm, k, 3, f, fmt=NUMF)
        put(sm, k, 4, f'=IF(ABS(C{k})<0.5,"OK","CHECK")', h="center")
    ls.status_formats(sm, "D18:D20", "D18")

    band(sm, 22, 2, 6, "Totals by FSLI (follows the TB sheet)")
    col_headers(sm, 23, ["", "FSLI category", "Lead ref", "CY balance", "PY balance", "Accounts"], start=1)
    r0 = 24
    k = r0
    for stmt, block, group, letter, its in SECTIONS:
        for f, _ in its:
            put(sm, k, 2, f)
            put(sm, k, 3, FSLI[f]["ref"], h="center")
            put(sm, k, 4, f"=SUMIFS({rng('F')},{rng('C')},$B{k})", fmt=NUMF)
            put(sm, k, 5, f"=SUMIFS({rng('E')},{rng('C')},$B{k})", fmt=NUMF)
            put(sm, k, 6, f"=COUNTIF({rng('C')},$B{k})", fmt="#,##0", h="center")
            k += 1
    put(sm, k, 2, "Not mapped / not recognised", italic=True)
    put(sm, k, 4, f"=TB!F{tr}-SUM(D{r0}:D{k - 1})", fmt=NUMF)
    put(sm, k, 5, f"=TB!E{tr}-SUM(E{r0}:E{k - 1})", fmt=NUMF)
    put(sm, k, 6, f"=C8-SUM(F{r0}:F{k - 1})", fmt="#,##0", h="center")
    put(sm, k + 1, 2, "Total (should be nil)", bold=True)
    put(sm, k + 1, 4, f"=SUM(D{r0}:D{k})", fmt=NUMF, bold=True, top=True)
    put(sm, k + 1, 5, f"=SUM(E{r0}:E{k})", fmt=NUMF, bold=True, top=True)
    put(sm, k + 1, 6, f"=SUM(F{r0}:F{k})", fmt="#,##0", bold=True, top=True, h="center")

    nr = k + 3
    band(sm, nr, 2, 6, "Next steps and how to read this")
    for j, t in enumerate([
        "1. Open the Review sheet and settle every Medium, Low and Unmapped account (change the FSLI Category on the TB sheet).",
        "2. Sign convention in the TB sheet: debit positive, credit negative. Adjustments are zero - fill them in if you have audit adjustments.",
        "3. Upload the TB sheet (save the file first) to the Lead Schedule tool. The extra columns H to J are ignored there.",
        "4. Next year, upload this file as the Previous mapping and unchanged accounts will map instantly.",
        "Suggestions are rule-based and are not a substitute for professional judgement: the auditor owns the final mapping."], 1):
        put(sm, nr + j, 2, t, italic=(j == 4), color=ls.GREY if j == 4 else None)
    widths(sm, {"A": 3, "B": 44, "C": 24, "D": 16, "E": 16, "F": 12})
    sm.sheet_view.showGridLines = False
    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)
    cnt = {g: sum(1 for x in rows if x["conf"] == g) for g in GRADE}
    return dict(n=n, counts=cnt, review=len(todo))


# --------------------------------------------------------------------------- #
# TEMPLATE (a raw client TB in Tally-like layout)
# --------------------------------------------------------------------------- #
SAMPLE = [  # code, name, group, CY Dr(+)/Cr(-)
    ("1001", "Freehold Land", "Fixed Assets", 42000), ("1002", "Factory Building", "Fixed Assets", 98000),
    ("1003", "Accumulated Depreciation - Building", "Fixed Assets", -29500), ("1004", "Plant & Machinery", "Fixed Assets", 210000),
    ("1005", "Acc. Depreciation - Plant & Machinery", "Fixed Assets", -114700), ("1006", "Computers & Laptops", "Fixed Assets", 14800),
    ("1007", "Capital Work-in-Progress", "Fixed Assets", 12400), ("1010", "Tally ERP Software", "Fixed Assets", 3400),
    ("1020", "Investment in Subsidiary - ABC Exports Pvt Ltd", "Investments", 15000), ("1021", "Axis Liquid Fund - Growth", "Investments", 7500),
    ("1030", "Security Deposit - Factory Landlord", "Deposits (Asset)", 2100),
    ("1100", "Raw Materials", "Stock-in-hand", 38600), ("1101", "Work in Progress", "Stock-in-hand", 14900),
    ("1102", "Finished Goods", "Stock-in-hand", 41200), ("1103", "Stores & Spares", "Stock-in-hand", 5300),
    ("1200", "Domestic Customers", "Sundry Debtors", 62400), ("1201", "Export Customers", "Sundry Debtors", 28700),
    ("1202", "Provision for Doubtful Debts", "Sundry Debtors", -3200),
    ("1300", "HDFC Bank - Current A/c", "Bank Accounts", 18400), ("1301", "SBI Collection A/c", "Bank Accounts", 6100),
    ("1302", "Fixed Deposit - ICICI", "Bank Accounts", 20000), ("1303", "Cash in Hand", "Cash-in-Hand", 140),
    ("1500", "GST Input Credit", "Duties & Taxes", 9300), ("1501", "Advances to Suppliers", "Loans & Advances (Asset)", 5600),
    ("1502", "Prepaid Insurance", "Loans & Advances (Asset)", 2700), ("1503", "TDS Receivable", "Duties & Taxes", 1800),
    ("1504", "Staff Advances", "Loans & Advances (Asset)", 950),
    ("1599", "Suspense A/c", "Suspense", 640),
    ("2001", "Equity Share Capital", "Capital Account", -50000), ("2002", "General Reserve", "Reserves & Surplus", -25000),
    ("2003", "Profit & Loss A/c (opening)", "Reserves & Surplus", None),
    ("2100", "HDFC Term Loan", "Secured Loans", -72000), ("2101", "Deferred Tax Liability", "Provisions", -9200),
    ("2102", "Provision for Gratuity (Non-current)", "Provisions", -7400), ("2103", "Loan from Director", "Unsecured Loans", -8000),
    ("2200", "Working Capital Demand Loan - HDFC", "Bank OD A/c", -34500), ("2201", "Current Maturities of Term Loan", "Secured Loans", -12000),
    ("2300", "Sundry Creditors - Materials", "Sundry Creditors", -44800), ("2301", "Sundry Creditors - Services", "Sundry Creditors", -9600),
    ("2302", "Creditors for Capital Goods", "Sundry Creditors", -4200),
    ("2400", "GST Payable", "Duties & Taxes", -6200), ("2401", "Salary Payable", "Current Liabilities", -5100),
    ("2402", "Advance from Customers", "Current Liabilities", -3800), ("2403", "Audit Fees Payable", "Current Liabilities", -4500),
    ("2500", "Provision for Gratuity (Current)", "Provisions", -1200), ("2501", "Provision for Income Tax", "Provisions", -11600),
    ("2502", "TDS Payable", "Duties & Taxes", -1450),
    ("3001", "Sales - Domestic", "Sales Accounts", -415000), ("3002", "Sales - Export", "Sales Accounts", -182000),
    ("3003", "Sales Returns", "Sales Accounts", 6200), ("3010", "Interest on Fixed Deposits", "Indirect Incomes", -1600),
    ("3011", "Rent Received", "Indirect Incomes", -2400), ("3012", "Export Incentive", "Indirect Incomes", -2900),
    ("4001", "Purchases - Raw Materials", "Purchase Accounts", 248000), ("4002", "Freight Inward", "Direct Expenses", 6800),
    ("4003", "Opening Stock - Finished Goods", "Direct Expenses", 36500), ("4004", "Wages - Factory", "Direct Expenses", 38200),
    ("4005", "Power & Fuel", "Direct Expenses", 14100),
    ("5001", "Salaries & Bonus", "Indirect Expenses", 52400), ("5002", "PF Employer Contribution", "Indirect Expenses", 4300),
    ("5003", "Interest on Term Loan", "Indirect Expenses", 8900), ("5004", "Bank Charges", "Indirect Expenses", 620),
    ("5005", "Depreciation", "Indirect Expenses", 19600), ("5006", "Rent - Office", "Indirect Expenses", 3600),
    ("5007", "Repairs & Maintenance - Building", "Indirect Expenses", 1900), ("5008", "Travelling & Conveyance", "Indirect Expenses", 2700),
    ("5009", "Professional Fees", "Indirect Expenses", 3100), ("5010", "Audit Fees", "Indirect Expenses", 900),
    ("5011", "Advertisement & Sales Promotion", "Indirect Expenses", 5200), ("5012", "Telephone & Internet", "Indirect Expenses", 840),
    ("5013", "Miscellaneous Expenses", "Indirect Expenses", 1240), ("5014", "Bad Debts Written Off", "Indirect Expenses", 780),
    ("5015", "Income Tax Expense", "Indirect Expenses", 11600), ("5016", "Sundry Items (unclassified)", "Indirect Expenses", 410),
]


def make_template(path):
    rows = [list(x) for x in SAMPLE]
    cy_total = sum(r[3] for r in rows if r[3] is not None)
    for r in rows:
        if r[3] is None:
            r[3] = -round(cy_total)                        # balancing figure so the demo TB nets to nil
    py = [round(r[3] * 0.9) for r in rows]
    bal = [i for i, r in enumerate(rows) if r[1].startswith("Profit & Loss")][0]
    py[bal] -= sum(py)
    wb = Workbook()
    ws = wb.active
    ws.title = "Instructions"
    lines = [("RAW TRIAL BALANCE - TB Mapping tool", True), ("", False),
             ("1.  Go to the 'TB' tab. Replace the sample rows with the client's trial balance (keep the header row).", False),
             ("2.  Needed: an account name and EITHER one balance column OR separate Debit and Credit columns.", False),
             ("3.  Optional but helpful: Account Code, Group (the ledger group / 'Under' in Tally), PY Balance.", False),
             ("4.  A single balance column is read as debit positive, credit negative (change this on the web form if yours is the other way round).", False),
             ("5.  Upload the file on the TB Mapping page. Optionally add last year's mapped TB so known accounts map identically.", False),
             ("", False), ("The sample data is a fictitious manufacturing company, with a few deliberately awkward ledgers - illustration only.", False)]
    for i, (t, b) in enumerate(lines, 1):
        put(ws, i, 1, t, bold=b, size=14 if b else 10, color=ls.NAVY if b else None)
    ws.column_dimensions["A"].width = 130
    tb = wb.create_sheet("TB")
    col_headers(tb, 1, ["Account Code", "Account Name", "Group", "PY Balance", "Debit", "Credit"])
    for i, (r, p) in enumerate(zip(rows, py), 2):
        put(tb, i, 1, r[0])
        put(tb, i, 2, r[1])
        put(tb, i, 3, r[2])
        put(tb, i, 4, p, fmt=NUMF, color=ls.BLUE)
        put(tb, i, 5, r[3] if r[3] > 0 else None, fmt=NUMF, color=ls.BLUE)
        put(tb, i, 6, -r[3] if r[3] < 0 else None, fmt=NUMF, color=ls.BLUE)
    widths(tb, {"A": 14, "B": 50, "C": 28, "D": 14, "E": 14, "F": 14})
    tb.freeze_panes = "A2"
    wb.save(path)
